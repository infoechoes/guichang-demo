"""Read cell values only. Never run macros, formulas, links, or workbook instructions."""
import io
import json
import sys
import zipfile
from pathlib import Path
import openpyxl


def read_xls(data):
    # The project-local reader is also included in the portable package.
    sys.path.insert(0, str(Path(__file__).parent / 'python-deps'))
    try:
        import xlrd
    except ImportError:
        raise ValueError('XLS读取引擎未安装，请安装local/excel-requirements.txt中的依赖')
    book = xlrd.open_workbook(file_contents=data, formatting_info=True,
                              on_demand=True, logfile=io.StringIO())
    try:
        if book.nsheets > 20:
            raise ValueError('最多支持20个工作表')
        sheets = []
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            if sheet.nrows > 5000 or sheet.ncols > 100:
                raise ValueError('单张工作表最多5000行、100列，请删除多余空行空列后上传')
            def value_at(row, col):
                cell = sheet.cell(row, col)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    return ''
                if cell.ctype == xlrd.XL_CELL_ERROR:
                    return xlrd.error_text_from_code.get(cell.value, '#ERROR!')
                if cell.ctype == xlrd.XL_CELL_DATE:
                    return xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat(sep=' ')
                if cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    return bool(cell.value)
                value = cell.value
                if isinstance(value, str):
                    return value[:2000]
                return value
            merged = {}
            for top, bottom, left, right in sheet.merged_cells:
                if bottom > 5000 or right > 100:
                    raise ValueError('合并单元格超过5000行或100列限制')
                value = value_at(top, left)
                for row in range(top, bottom):
                    for col in range(left, right):
                        merged[row, col] = value
            rows = []
            for row in range(sheet.nrows):
                values = [merged.get((row, col), value_at(row, col)) for col in range(sheet.ncols)]
                if any(value != '' for value in values):
                    rows.append({'number': row + 1, 'cells': values})
            sheets.append({'name': sheet.name, 'hidden': bool(sheet.visibility), 'rows': rows})
        return sheets
    finally:
        book.release_resources()


try:
    data = sys.stdin.buffer.read(10 * 1024 * 1024 + 1)
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('Excel不能超过10MB')
    # Some genuine OLE/XLS files embed ZIP metadata; is_zipfile alone misclassifies them.
    if not data.startswith(b'PK\x03\x04'):
        print(json.dumps({'sheets': read_xls(data)}, ensure_ascii=True, allow_nan=False))
        sys.exit(0)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
            raise ValueError('表格解压后过大，请只保留需要比价的工作表')
        if any('vbaProject' in i.filename for i in z.infolist()):
            raise ValueError('不支持带宏的工作簿，请另存为xlsx')
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False, keep_links=False)
    if len(wb.worksheets) > 20:
        raise ValueError('最多支持20个工作表')
    sheets = []
    for sheet in wb.worksheets:
        if sheet.max_row > 5000 or sheet.max_column > 100:
            raise ValueError('单张工作表最多5000行、100列，请删除多余空行空列后上传')
        merged = {}
        for area in sheet.merged_cells.ranges:
            value = sheet.cell(area.min_row, area.min_col).value
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    merged[row, col] = value
        rows = []
        for row in sheet.iter_rows():
            values = []
            for cell in row:
                v = merged.get((cell.row, cell.column), cell.value)
                if v is None:
                    v = ''
                elif not isinstance(v, (str, int, float, bool)):
                    v = str(v)
                if isinstance(v, str) and len(v) > 2000:
                    v = v[:2000]
                values.append(v)
            if any(v != '' for v in values):
                rows.append({'number': row[0].row, 'cells': values})
        sheets.append({'name': sheet.title, 'hidden': sheet.sheet_state != 'visible', 'rows': rows})
    print(json.dumps({'sheets': sheets}, ensure_ascii=True, allow_nan=False))
except Exception as e:
    message = str(e) if isinstance(e, ValueError) else '无法读取Excel，请确认文件为真实的xls/xlsx工作簿，且未加密、未损坏'
    print(json.dumps({'error': message}, ensure_ascii=True))
    sys.exit(1)
