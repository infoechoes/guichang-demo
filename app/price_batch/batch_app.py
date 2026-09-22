import io, json, uuid, datetime, zipfile, threading, copy
from pathlib import Path
from http.server import ThreadingHTTPServer
from urllib.parse import urlsplit, parse_qs
from decimal import Decimal
from .pricing import projects, save, calculate, number
from html import escape
from .price_handoff import identity, current_result, handoff
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from .xinfadi_server import Handler, lookup

ROOT=Path(__file__).parent
SESSIONS={}
EXTRA=['匹配品名','最低价','平均价','最高价','行情单位','实际行情日期','来源','查询状态','单位换算比例','行情说明','项目','价格用途','取价基准','加工费','浮动数值','计算式','最终参考价','规则版本','来源批次ID','来源行ID','输入指纹','计算批次ID']

def load(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if sum(i.file_size for i in z.infolist())>60_000_000: raise ValueError('Excel 解压后过大（上限 60 MB）')
    return openpyxl.load_workbook(io.BytesIO(raw))

def preview(s, sheet, header):
    wb=load(s['raw']); ws=wb[sheet]
    if ws.max_row>10000 or ws.max_column>200: raise ValueError('支持最多 10000 行、200 列')
    if not 1<=header<=ws.max_row: raise ValueError('表头行无效')
    s.update(sheet=sheet,header=header,results={},mapping=None,project=None,calculation_batch_id=None)
    return {'sheets':wb.sheetnames,'sheet':sheet,'header':header,'headers':[str(c.value or '') for c in ws[header]],'rows':[[c.value for c in row] for row in ws.iter_rows(min_row=header+1)],'count':ws.max_row-header,'source_batch_id':identity(s,ws,header)['source_batch_id'],'row_identities':[identity(s,ws,row) for row in range(header+1,ws.max_row+1)]}

def celltext(v): return str(v).strip() if v is not None else ''
def sameunit(a,b):
    aliases={'公斤':'kg','千克':'kg','KG':'kg','市斤':'斤','克':'g'}
    a,b=celltext(a),celltext(b)
    return bool(a and b) and aliases.get(a,a)==aliases.get(b,b)

def confirm(s,row,index,factor,fee=None,adjustment=None):
    ws=load(s['raw'])[s['sheet']]
    r=current_result(s,ws,row)
    if not r or not 0<=index<len(r['items']): raise ValueError('请先查询并选择有效候选')
    if r.get('protected'): raise ValueError('原单已有价格，保留且不浮动')
    rule=s.get('project')
    if not rule: raise ValueError('请先保存并选择项目规则，再查询')
    item=r['items'][index]
    if not item.get('published_at') or not item.get('reference_url'):raise ValueError('候选缺少日期或来源，不能确认')
    ratio=number(factor)
    if sameunit(r['unit'],item['unit']) and ratio!=1: raise ValueError('单位一致时换算比例必须为 1')
    base=item.get(rule['basis'])
    if base is None:raise ValueError('上游缺少项目指定的基准价格，不能填价')
    result=calculate(base,ratio,rule,fee,adjustment)
    r.update(result,selected=index,status='已确认参考价（待业务结算核对）',confirmed=True)
    return r

def print_records(s,rows):
    if not rows or len(rows)>20 or len(set(rows))!=len(rows):raise ValueError('请选择1至20条不同商品')
    records=[]
    for row in rows:
        r=current_result(s,load(s['raw'])[s['sheet']],int(row)) or {}
        if not r.get('confirmed'):raise ValueError('勾选行尚未确认，不能打印凭证')
        item=r['items'][r['selected']] if r.get('selected') is not None else {}
        records.append({'row':row,'name':r.get('name'),'unit':r.get('unit'),'price':r.get('price_decimal',r.get('original_price')),'status':r['status'],'basis':r.get('rule',{}).get('basis','原单已有价'),'base_price':r.get('base_price','—'),'formula':r.get('formula','保留原价，不再浮动'),'source':item.get('reference_url') or r.get('source','原始上传表'),'date':item.get('published_at','')[:10] or r.get('date','未提供'),'matched_name':item.get('name','—'),'spec':item.get('spec',''),'origin':item.get('origin',''),'market_unit':item.get('unit','')})
    return {'project':s['project'],'records':records,'generated_at':datetime.datetime.now().isoformat(timespec='seconds'),'evidence_type':'自主生成价格汇总，非网站截图；非客户结算确认'}

def print_html(data):
    e=lambda v:escape(str(v))
    trs=''.join('<tr>'+''.join('<td>'+e(v)+'</td>' for v in [r['row'],r['name'],r['matched_name']+' '+r['spec']+' '+r['origin'],r['date'],r['base_price']+'/'+r['market_unit'],r['formula'].replace('ROUNDDOWN','截2位').replace('ROUND','四舍五入2位').replace('none','不舍入'),str(r['price'])+'/'+r['unit']])+'</tr>' for r in data['records'])
    sources=list(dict.fromkeys(r['source'] for r in data['records']))
    return '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>价格核对凭证</title><style>@page{size:A4 portrait;margin:10mm}body{font:10px/1.25 sans-serif;color:#17271d}table{width:100%;border-collapse:collapse;table-layout:fixed}td,th{border:1px solid #aab4ab;padding:3px;overflow-wrap:anywhere}tr{break-inside:avoid}th:nth-child(1){width:4%}th:nth-child(2){width:13%}th:nth-child(3){width:15%}th:nth-child(4){width:13%}th:nth-child(5){width:12%}th:nth-child(6){width:27%}th:nth-child(7){width:16%}h1{font-size:19px}footer{font-size:10px;overflow-wrap:anywhere}button{margin:12px}@media print{button{display:none}}</style><button onclick="window.print()">打印 / 保存为 PDF</button><h1>'+e(data['project']['name'])+' · 价格核对凭证</h1><p>'+e('供应商采购协议参考价' if data['project']['purpose']=='purchase' else '客户销售参考价')+' · 规则版本 '+e(data['project']['version'])+' · '+e(data['generated_at'])+'</p><p>'+e(data['evidence_type'])+'</p><table><thead><tr><th>行</th><th>原品名</th><th>候选 / 规格 / 产地</th><th>行情日期</th><th>基准价/单位</th><th>计算规则</th><th>最终参考价</th></tr></thead><tbody>'+trs+'</tbody></table><footer><p>基准：'+e({'average_price':'平均价','low_price':'最低价','high_price':'最高价','fixed':'固定价'}[data['project']['basis']])+'；加工费按原单单位计；ROUND 四舍五入两位，ROUNDDOWN 截两位，none 不舍入。</p><p>来源：'+e('；'.join(sources))+'</p><p>保留原价的行不代表本次重新查证。真实网页证据请从来源链接打开原站另行核验。</p></footer></html>'

def export(s):
    wb=load(s['raw']); ws=wb[s['sheet']]; source_ws=load(s['raw'])[s['sheet']]; h=s['header']; base=ws.max_column
    existing=[i for i,c in enumerate(ws[h],1) if celltext(c.value)=='单价']
    pricecol=existing[0] if len(existing)==1 else base+1
    start=base+1 if existing and len(existing)==1 else base+2
    if pricecol>base: ws.cell(h,pricecol,'单价')
    for i,label in enumerate(EXTRA,start): ws.cell(h,i,label)
    for row in range(h+1,ws.max_row+1):
        r=current_result(s,source_ws,row) or {'status':'未查询或输入已变化，待核','items':[]}
        item=r['items'][r['selected']] if r.get('selected') is not None else {}
        status=r['status']
        if r.get('price') is not None:
            if ws.cell(row,pricecol).value is None: ws.cell(row,pricecol,r['price'])
            else: status+='；原单价非空，未覆盖'
        vals=[item.get('name'),item.get('low_price'),item.get('average_price'),item.get('high_price'),item.get('unit'),(item.get('published_at') or '')[:10] or None,item.get('reference_url'),status,r.get('factor'),'自主生成参考价，非客户结算价',(s.get('project') or {}).get('name'),(s.get('project') or {}).get('purpose'),r.get('rule',{}).get('basis'),r.get('fee'),r.get('adjustment'),r.get('formula'),r.get('price_decimal'),r.get('rule',{}).get('version'),identity(s,source_ws,row)['source_batch_id'],identity(s,source_ws,row)['source_line_id'],identity(s,source_ws,row)['input_fingerprint'],s.get('calculation_batch_id')]
        for offset,value in enumerate(vals):
            c=ws.cell(row,start+offset)
            c.value=float(value) if offset in (1,2,3,8) and value is not None else value
            if isinstance(c.value,str): c.data_type='s'
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if offset in (1,2,3): c.number_format='0.00##'
        ws.cell(row,pricecol).number_format='0.00##'
    for col in range(start,start+len(EXTRA)):
        c=ws.cell(h,col); c.fill=PatternFill('solid',fgColor='174B3C'); c.font=Font(color='FFFFFF',bold=True)
        ws.column_dimensions[c.column_letter].width=24 if col!=start+6 else 48
    out=io.BytesIO(); wb.save(out); return out.getvalue()

class App(Handler):
    def respond(self,status,value,content_type='application/json; charset=utf-8'):
        if not isinstance(value,(str,bytes)): value=json.dumps(value,ensure_ascii=False,default=str)
        raw=value if isinstance(value,bytes) else value.encode()
        self.send_response(status); self.send_header('Content-Type',content_type); self.send_header('Content-Length',str(len(raw))); self.send_header('Cache-Control','no-store'); self.send_header('X-Content-Type-Options','nosniff'); self.end_headers(); self.wfile.write(raw)
    def do_GET(self):
        p=urlsplit(self.path)
        if p.path in ('/','/app.js','/style.css'):
            name={'/':'index.html','/app.js':'app.js','/style.css':'style.css'}[p.path]
            return self.respond(200,(ROOT/name).read_bytes(),{'/':'text/html; charset=utf-8','/app.js':'application/javascript; charset=utf-8','/style.css':'text/css; charset=utf-8'}[p.path])
        if p.path=='/api/projects': return self.respond(200,projects())
        if p.path=='/health': return self.respond(200,{'status':'ok','local_only':True})
        return self.respond(404,{'error':'不存在'})
    def do_POST(self):
        try:
            origin=self.headers.get('Origin')
            if origin and origin!='http://'+self.headers.get('Host',''): raise ValueError('不接受跨站请求')
            size=int(self.headers.get('Content-Length','0'))
            if not 0<size<=15_000_000: raise ValueError('文件/请求大小上限 15 MB')
            raw=self.rfile.read(size); path=urlsplit(self.path).path
            if path=='/api/upload':
                wb=load(raw); sid=uuid.uuid4().hex
                if len(SESSIONS)>=20: SESSIONS.pop(next(iter(SESSIONS)))
                s={'raw':raw}; result=preview(s,wb.active.title,1); SESSIONS[sid]=s
                return self.respond(200,dict(result,id=sid))
            d=json.loads(raw)
            if path=='/api/projects': return self.respond(200,save(d['project']))
            s=SESSIONS.get(d.get('id'))
            if not s: raise ValueError('会话已失效，请重新上传')
            if path=='/api/preview': return self.respond(200,preview(s,d['sheet'],int(d['header'])))
            if path=='/api/setup':
                wb=load(s['raw']); width=wb[s['sheet']].max_column
                mapping=d['mapping']
                for key in ('name','spec','unit','quantity'):
                    c=int(mapping[key]);
                    if not -1<=c<width or (key=='name' and c<0): raise ValueError('品名列必须指定，其他列可留空')
                project=projects().get(d.get('project_id'))
                if d.get('project_id') and not project:raise ValueError('所选项目不存在')
                s.update(mapping=mapping,date=d['date'],results={},project=copy.deepcopy(project),calculation_batch_id=uuid.uuid4().hex); return self.respond(200,{'ok':True,'calculation_batch_id':s['calculation_batch_id']})
            if path=='/api/query':
                if not s.get('mapping'): raise ValueError('请先确认列与日期')
                if not s.get('project'):raise ValueError('查价前请先保存并选择项目规则；未定价订单仍可完整交接为待核')
                row=int(d['row']); ws=load(s['raw'])[s['sheet']]
                if not s['header']<row<=ws.max_row: raise ValueError('行号超出范围')
                m=s['mapping']; name=celltext(ws.cell(row,int(m['name'])+1).value); unit=celltext(ws.cell(row,int(m['unit'])+1).value) if int(m['unit'])>=0 else ''
                r={**identity(s,ws,row),'items':[],'status':'品名为空，未查询','unit':unit,'name':name}
                price_cols=[c.column for c in ws[s['header']] if celltext(c.value) in ('单价','固定价','协议价')]
                existing=next((ws.cell(row,c).value for c in price_cols if ws.cell(row,c).value is not None),None)
                if existing is not None:
                    r.update(protected=True,original_price=existing,status='原单已有价，保留且不浮动',confirmed=True,source='原上传表 '+s['sheet']+' 第'+str(row)+'行',date='原单日期未核')
                    s['results'][row]=r;return self.respond(200,r)
                if s['project']['basis']=='fixed' and name:
                    rule=s['project'];r.update(price=float(number(rule['fixed_price'])),price_decimal=rule['fixed_price'],rule=rule,status='项目固定参考价，未浮动',confirmed=True,formula='项目固定价，未浮动',source='本地保存项目规则 v'+str(rule['version']),date='固定价有效日期待业务核对')
                    s['results'][row]=r;return self.respond(200,r)
                if name:
                    try:
                        data=lookup(name,s['date']); r.update(items=data['items'],status='待选择并确认匹配' if data['items'] else '指定日期范围无同名行情，价格留空')
                    except Exception as e: r['status']='请求失败：'+str(e)+'；价格留空'
                if r['calculation_batch_id']!=s.get('calculation_batch_id'):raise ValueError('查价期间批次设置变化，请重新查询')
                s['results'][row]=r; return self.respond(200,r)
            if path=='/api/confirm':
                if d.get('calculation_batch_id') is not None and d['calculation_batch_id']!=s.get('calculation_batch_id'):raise ValueError('计算批次已变化，请重新查询')
                if d.get('source_line_id') is not None and d['source_line_id']!=identity(s,load(s['raw'])[s['sheet']],int(d['row']))['source_line_id']:raise ValueError('来源行不一致')
                return self.respond(200,confirm(s,int(d['row']),int(d['index']),d['factor'],d.get('fee'),d.get('adjustment')))
            if path=='/api/clear':
                r=s['results'][int(d['row'])]
                if r.get('protected'):raise ValueError('原单价格不可撤销')
                for k in ('selected','price','price_decimal','factor','confirmed','formula','rule','fee','adjustment','base_price'): r.pop(k,None)
                r['status']='待选择并确认匹配'; return self.respond(200,r)
            if path=='/api/print': return self.respond(200,{'html':print_html(print_records(s,d['rows']))})
            if path=='/api/handoff':
                return self.respond(200,handoff(s,load(s['raw'])[s['sheet']],d.get('rows'),d.get('line_ids')))
            if path=='/api/export': return self.respond(200,export(s),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            return self.respond(404,{'error':'不存在'})
        except Exception as e: return self.respond(400,{'error':str(e)})

if __name__=='__main__':
    print('本机查价页面：http://127.0.0.1:42871',flush=True)
    ThreadingHTTPServer(('127.0.0.1',42871),App).serve_forever()
