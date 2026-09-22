"""Shared tax knowledge and account-scoped, fresh-review invoice preparation."""
import base64
import csv
import datetime as dt
import io
import json
import os
import re
import shutil
import threading
import uuid
import zipfile
from decimal import Decimal
from pathlib import Path
import openpyxl
from .library import Library, digest, file_hash, js, stamp, text, rate, candidate_options, date_clause
from .tax_template import inspect_template, write_template, number

FIELDS={'name':['品名','商品名称','名称','项目名称'],'spec':['规格型号','规格'],'unit':['单位','计量单位'],'qty':['计划采购','商品数量','数量'],'unit_price':['商品单价','单价'],'amount':['总价','金额','价税合计'],'tax_code':['商品和服务税收分类编码','税收分类编码'],'rate':['税率']}
REASONS=('商品不对','编码不对','税率不对','信息不足')
REVIEW_FIELDS=('name','tax_code','rate','spec','unit','identity','facts','process','seller','taxpayer','trade','valid_from','valid_to','basis','treatment','policy','export_policy')
RISK_SHA='d4db2dad09458d0baa07b141e9b962dff496444b2b0a211d819e3bc8b57978ff'

def uid():return uuid.uuid4().hex
def assert_version(record,p):
    if record['version']!=p.get('version'):raise ValueError('版本过期，请重新读取；未覆盖最新修改')
def review_value(value):
    if not isinstance(value,dict):raise ValueError('复核字段格式无效')
    r={k:text(value.get(k,'')) for k in REVIEW_FIELDS}
    r['rate']=rate(r['rate'])
    r['same_product']=value.get('same_product') is True
    r['conditions_checked']=value.get('conditions_checked') is True
    r['conflict_ack']=value.get('conflict_ack') is True
    if any(len(v)>4000 for v in r.values() if isinstance(v,str)):raise ValueError('单个复核字段过长')
    return r
def validate_yes(r,conflict=False):
    required={'seller':'适用销售主体','name':'真实同品名称','tax_code':'税收分类码','unit':'单位','identity':'同一真实商品的识别依据','facts':'商品事实/规格核实','process':'加工及包装事实','seller':'适用销售主体','taxpayer':'纳税/计税身份','trade':'交易适用条件','valid_from':'适用起日','valid_to':'适用止日','basis':'理由及依据'}
    missing=[label for k,label in required.items() if not r.get(k)]
    if r.get('rate') not in ('0','0.09','0.13'):missing.append('0/9%/13%单选；其他历史值须单独待核')
    if not re.fullmatch(r'\d{19}',r.get('tax_code','')):missing.append('19位文本税收分类码')
    if not r.get('same_product'):missing.append('确认未替换实际交易商品')
    if not r.get('conditions_checked'):missing.append('适用条件已核实')
    if conflict and not r.get('conflict_ack'):missing.append('明确处理历史或知识冲突')
    if r.get('rate')=='0':
        if r.get('treatment') not in ('免税','零税率'):missing.append('0对应免税还是零税率')
        if not r.get('policy'):missing.append('0对应政策及成立条件')
        if r.get('treatment')=='免税' and r.get('export_policy')!='免税':missing.append('免税对应模板优惠政策必须为免税')
        if r.get('treatment')=='零税率' and r.get('export_policy'):missing.append('零税率不等同免税，模板优惠政策须留空并保留政策依据')
    elif r.get('treatment') not in ('','正常征税') or r.get('export_policy'):
        missing.append('9%/13%不能沿用0待遇或未经核实的优惠枚举')
    try:
        start=dt.date.fromisoformat(r['valid_from']);end=dt.date.fromisoformat(r['valid_to'])
        if start>end:missing.append('有效适用期间')
    except (ValueError,KeyError):missing.append('有效适用日期')
    if missing:raise ValueError('暂不能通过：'+'；'.join(missing))

class Workflow:
    def __init__(self,db,owner,private_root,actor=None):
        self.lib=Library(db);self.owner=owner;self.actor=actor or owner
        self.root=Path(private_root);self.root.mkdir(parents=True,exist_ok=True)
        self.files=self.root/'v2';self.files.mkdir(exist_ok=True)
    def batch(self,c,bid):
        b=c.execute('SELECT * FROM batches WHERE id=? AND owner=?',(bid,self.owner)).fetchone()
        if not b:raise ValueError('当前账号无此批次')
        b=dict(b);b['context']=json.loads(b['context']);return b
    def item(self,c,iid):
        r=c.execute('SELECT i.* FROM items i JOIN batches b ON b.id=i.batch_id WHERE i.id=? AND b.owner=?',(iid,self.owner)).fetchone()
        if not r:raise ValueError('当前账号无此行')
        r=dict(r);r['original']=json.loads(r['original']);r['review']=json.loads(r['review']);return r
    def upload_start(self,p):
        name=Path(text(p.get('name')).replace('\\','/')).name
        if Path(name).suffix.lower() not in ('.xlsx','.csv'):raise ValueError('支持xlsx或csv')
        token=uid();path=self.files/(token+Path(name).suffix.lower());path.touch(mode=0o600)
        with self.lib.connect() as c:c.execute('INSERT INTO uploads(id,owner,name,path) VALUES(?,?,?,?)',(token,self.owner,name,str(path)))
        return {'token':token,'offset':0}
    def uploaded(self,c,token,complete=True):
        r=c.execute('SELECT * FROM uploads WHERE id=? AND owner=?',(token,self.owner)).fetchone()
        if not r or (complete and not r['complete']):raise ValueError('当前账号无此完整上传')
        return dict(r)
    def upload_chunk(self,p):
        raw=base64.b64decode(p.get('data',''),validate=True)
        if len(raw)>2*1024*1024:raise ValueError('每块最多2MB')
        with self.lib.connect() as c:
            c.execute('BEGIN IMMEDIATE');u=self.uploaded(c,p['token'],False)
            if u['complete'] or p.get('offset')!=u['size']:raise ValueError('上传分块偏移不符，请重试当前块')
            if u['size']+len(raw)>200*1024*1024:raise ValueError('文件超过200MB')
            with open(u['path'],'r+b') as f:
                f.seek(u['size']);f.write(raw);f.truncate();f.flush()
            size=u['size']+len(raw);done=p.get('done') is True
            sha=file_hash(u['path']) if done else ''
            c.execute('UPDATE uploads SET size=?,complete=?,sha=? WHERE id=?',(size,int(done),sha,u['id']))
        return {'token':u['id'],'offset':size,'complete':done,'sha256':sha}
    def workbook(self,path):
        with zipfile.ZipFile(path) as z:
            if len(z.infolist())>5000 or sum(i.file_size for i in z.infolist())>1024*1024*1024:raise ValueError('工作簿解压结构过大')
        return openpyxl.load_workbook(path,read_only=True,data_only=True)
    def rows(self,u,sheet):
        if u['name'].lower().endswith('.csv'):
            # Iterative text IO, not a whole-sheet list.
            f=open(u['path'],encoding='utf-8-sig',newline='')
            try:yield from csv.reader(f)
            finally:f.close()
        else:
            w=self.workbook(u['path'])
            try:
                if sheet not in w.sheetnames:raise ValueError('工作表不存在')
                yield from w[sheet].iter_rows(values_only=True)
            finally:w.close()
    def upload_info(self,p):
        with self.lib.connect() as c:u=self.uploaded(c,p['token'])
        if u['name'].lower().endswith('.csv'):sheets=[{'name':'CSV'}]
        else:
            w=self.workbook(u['path']);sheets=[{'name':s.title,'rows':s.max_row,'cols':s.max_column} for s in w];w.close()
        sheet=p.get('sheet',sheets[0]['name']);header=int(p.get('header',0))
        if not header:
            choices=[];scan=self.rows(u,sheet)
            try:
                for n,row in enumerate(scan,1):
                    values={text(x) for x in row}
                    choices.append((sum(any(a in values for a in aliases) for aliases in FIELDS.values()),-n))
                    if n>=15:break
            finally:scan.close()
            header=-max(choices)[1] if choices else 1
        if not 1<=header<=100:raise ValueError('表头行须在1至100')
        head=[];preview=[];pre=[]
        iterator=self.rows(u,sheet)
        try:
            for n,row in enumerate(iterator,1):
                if n<header:pre.extend(text(x) for x in row)
                if n==header:head=[text(v) for v in row]
                elif header<n<=header+5:preview.append([text(v) for v in row])
                if n>=header+5:break
        finally:iterator.close()
        return {'token':u['id'],'name':u['name'],'sha256':u['sha'],'sheets':sheets,'sheet':sheet,'header':header,'headers':head,'mapping':{k:next((i for i,h in enumerate(head) if h in aliases),-1) for k,aliases in FIELDS.items()},'preview':preview,'risk':self.risk(u,pre)}
    def risk(self,u,pre):
        if u['sha']==RISK_SHA or '示例客户C8月开票明细' in u['name'] or re.search('替换物品开票|换品开票|更换品名开票',' '.join(pre)):
            return '原件含“替换物品开票”风险。本批仅供结构核对；不得通过改名、人工通过或Agent建议生成可开票成品。请先核实真实交易。'
        return ''
    def start_history(self,p):
        with self.lib.connect() as c:u=self.uploaded(c,p['token'])
        if not u['name'].endswith('.xlsx'):raise ValueError('历史大表需xlsx')
        with self.lib.connect() as c:
            old=c.execute('SELECT * FROM imports WHERE sha=?',(u['sha'],)).fetchone()
            if old and old['status'] in ('running','complete'):return dict(old)
            c.execute("INSERT INTO imports(sha,name,kind,status,created) VALUES(?,?,'history','running',?) ON CONFLICT(sha) DO UPDATE SET status='running',error=''",(u['sha'],u['name'],stamp()))
        def run():
            try:self.lib.import_history(u['path'],u['sha'])
            except Exception:pass # Durable error is returned by /stats.
            finally:
                with self.lib.connect() as c:c.execute('UPDATE imports SET name=? WHERE sha=?',(u['name'],u['sha']))
        threading.Thread(target=run,daemon=True).start()
        return {'sha':u['sha'],'status':'running'}
    def import_batch(self,p):
        with self.lib.connect() as c:u=self.uploaded(c,p['token'])
        info=self.upload_info(p);m=p['mapping'];indexes=[int(v) for v in m.values() if int(v)>=0]
        if len(set(indexes))!=len(indexes) or any(i>=len(info['headers']) for i in indexes):raise ValueError('列映射重复或越界')
        if any(int(m.get(k,-1))<0 for k in ('name','qty','amount')):raise ValueError('请映射品名、数量、金额列；缺值仍须逐行待核')
        bid=uid();header=info['header'];nitems=0;amount=Decimal('0');pre=[]
        with self.lib.connect() as c:
            c.execute('INSERT INTO batches(id,owner,title,is_test,source_sha,created) VALUES(?,?,?,?,?,?)',(bid,self.owner,u['name']+' · '+info['sheet'],int(p.get('is_test') is True or os.environ.get('GC_TAX_TEST_MODE')=='1'),u['sha'],stamp()))
            for n,row in enumerate(self.rows(u,info['sheet']),1):
                if n<header:pre.extend(text(x) for x in row);continue
                if n==header:continue
                def get(k):
                    i=int(m.get(k,-1));return text(row[i]) if 0<=i<len(row) else ''
                original={k:get(k) for k in FIELDS}
                if not any(original.values()):continue
                if not original['name'] and any(re.search('合计|总计',text(x)) for x in row[:3]):continue
                if original['name'] in ('合计','总计'):continue
                nitems+=1
                original.update(source_sheet=info['sheet'],formula_note='使用原文件缓存数值；空缓存保持缺失，不重算、不换算金额')
                if original['amount']:
                    try:amount+=number(original['amount'],'金额')
                    except ValueError:pass
                c.execute('INSERT INTO items(id,batch_id,ordinal,source_row,original) VALUES(?,?,?,?,?)',(uid(),bid,nitems,n,js(original)))
            if not nitems:raise ValueError('没有业务明细')
            c.execute('UPDATE batches SET risk=? WHERE id=?',(self.risk(u,pre),bid))
            self.lib.audit(c,'batch',bid,self.actor,'导入新批次，所有行须新鲜复核',{}, {'count':nitems,'amount_sum':str(amount),'sha':u['sha']})
        return self.batch_view({'id':bid})
    def batch_view(self,p):
        offset=max(0,int(p.get('offset',0)));limit=max(1,min(int(p.get('limit',30)),100))
        with self.lib.connect() as c:
            b=self.batch(c,p['id']);states={r[0]:r[1] for r in c.execute('SELECT status,COUNT(*) FROM items WHERE batch_id=? GROUP BY status',(b['id'],))}
            b['total']=sum(states.values());b['progress']=states;b['offset']=offset;b['limit']=limit
            options=candidate_options(p);view=p.get('view','')
            if view not in ('','reusable','needs_check','confirmed'):raise ValueError('清单筛选无效')
            all_rows=[self.item(c,r[0]) for r in c.execute('SELECT id FROM items WHERE batch_id=? ORDER BY ordinal',(b['id'],))]
            summary={'reusable':0,'needs_check':0,'confirmed':0}
            for r in all_rows:
                r['candidates']=self.candidates(c,r['original'],b['context'],options)
                r['match_state']='confirmed' if r['status']=='yes' else ('reusable' if r['status']=='pending' and any(x['applicable'] for x in r['candidates']) and not b['risk'] else 'needs_check')
                summary[r['match_state']]+=1
            selected=[r for r in all_rows if not view or r['match_state']==view]
            b['match_summary']=summary;b['filtered_total']=len(selected);b['view']=view
            b['rows']=selected[offset:offset+limit]
            b['ready']=not b['risk'] and states.get('yes',0)==b['total'] and b['context'].get('basis_confirmed') is True
            return b
    def candidates(self,c,o,ctx,options=None):
        options=candidate_options(options or {});dc,da=date_clause(options)
        # Compute applicability and disagreements before limiting or filtering history.
        ids=[x[0] for x in c.execute("SELECT id FROM groups WHERE name=? OR (json_extract(review,'$.name')=? AND status='yes')",(o['name'],o['name']))]
        out=[]
        for gid in ids:
            g=self.lib.group(gid,c);r=g['review']
            exact=bool(o.get('spec') and o.get('unit')) and r.get('spec',g['spec'])==o.get('spec') and r.get('unit',g['unit'])==o.get('unit')
            applicable=exact and not r.get('annotation_only') and g['status']=='yes' and (not g['conflict'] or r.get('conflict_ack')) and all(ctx.get(k) and ctx.get(k)==r.get(k) for k in ('seller','taxpayer','trade','process')) and r.get('valid_from','~')<=ctx.get('date','')<=r.get('valid_to','')
            if applicable:
                try:validate_yes(r,bool(g['conflict']))
                except ValueError:applicable=False
            dated=dict(c.execute('SELECT COUNT(*) occurrences,MAX(d.source_date) last_used FROM observations o LEFT JOIN observation_dates d ON d.id=o.id WHERE o.group_id=? AND '+dc,[gid]+da).fetchone())
            out.append({**{k:g[k] for k in ('id','version','name','spec','unit','history','status','conflict','flags','source_files','confirmation','unknown_dates','count_basis','date_basis')},**dated,'review':r if g['status']=='yes' else {},'applicable':applicable,'exact':exact,'in_range':dc=='1=1' or dated['occurrences']>0,'reason':'已确认且同品规格/单位/交易条件/日期一致，可复用字段；仍须本行人工通过' if applicable else '仅候选：商品、主体、日期或适用条件待核，不自动套用'})
        approved=[g for g in out if g['applicable']]
        signatures={tuple(g['review'].get(k,'') for k in ('name','tax_code','rate','treatment','policy','export_policy')) for g in approved}
        if len(signatures)>1:
            for g in approved:g.update(applicable=False,conflict=True,reason='同条件存在多个不同确认结果，须处理差异；不按近期或次数自动选择')
        # Never hide a confirmed record or disagreement via historical date filters.
        out=[g for g in out if g['in_range'] or g['status']=='yes' or g['conflict']]
        out.sort(key=lambda g:(not g['applicable'],not g['exact'],g['id']))
        if options['sort']:
            key='last_used' if options['sort']=='recent' else 'occurrences'
            out.sort(key=lambda g:g[key] or ('' if key=='last_used' else 0),reverse=True)
            out.sort(key=lambda g:not g['applicable'])
        return out[:10]
    def save_context(self,p):
        ctx={k:text(p.get('context',{}).get(k,'')) for k in ('seller','taxpayer','trade','process','date','amount_basis','basis_evidence','reviewer')}
        ctx['basis_confirmed']=p.get('context',{}).get('basis_confirmed') is True
        if ctx['basis_confirmed'] and (ctx['amount_basis'] not in ('含税','不含税') or not ctx['basis_evidence']):raise ValueError('确认含税口径时需选择含税/不含税并填写依据，不自动换算')
        with self.lib.connect() as c:
            c.execute('BEGIN IMMEDIATE');b=self.batch(c,p['id']);assert_version(b,p)
            if b['context']!=ctx:
                self.lib.audit(c,'batch',b['id'],self.actor,'共同条件修改，所有通过行重审',b['context'],ctx)
                c.execute("UPDATE items SET status=CASE WHEN status='yes' THEN 'pending' ELSE status END,version=version+1 WHERE batch_id=?",(b['id'],))
                c.execute('UPDATE batches SET context=?,version=version+1 WHERE id=?',(js(ctx),b['id']))
        return self.batch_view({'id':p['id']})
    def review(self,p):
        kind=p.get('kind','group');action=p.get('action');target=p['id']
        if kind not in ('group','item') or action not in ('edit','yes','no'):raise ValueError('复核操作无效')
        with self.lib.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            obj=self.lib.group(target,c) if kind=='group' else self.item(c,target);assert_version(obj,p)
            before={'review':obj['review'],'status':obj['status'],'version':obj['version'],'reason':obj['reason']}
            reason='';r=obj['review'];status='pending';warning=''
            if action=='edit':r=review_value(p.get('review',{}))
            elif action=='no':
                category=p.get('reason_type');note=text(p.get('reason'))
                if category not in REASONS or not note:raise ValueError('No必须选择理由并填写说明；原记录不会删除')
                reason=category+'：'+note;status='no'
            else:
                if not r:raise ValueError('请先保存修改，再显式通过')
                if p.get('seen') is not True:raise ValueError('每条都须本次人工确认，不支持批量自动通过')
                conflict=bool(obj.get('conflict'))
                b=None
                if kind=='item':
                    b=self.batch(c,obj['batch_id']);ctx=b['context']
                    if b['risk']:raise ValueError(b['risk'])
                    if not all(ctx.get(k) for k in ('seller','taxpayer','trade','process','date','reviewer')) or not ctx.get('basis_confirmed'):raise ValueError('先核实共同条件、交易日期、含税口径和复核人')
                    if any(r.get(k)!=ctx[k] for k in ('seller','taxpayer','trade','process')):raise ValueError('行适用条件与批次条件不一致，请修改后重审')
                    if not r.get('valid_from','~')<=ctx['date']<=r.get('valid_to',''):raise ValueError('适用期间未覆盖本次交易日期')
                    original=obj['original'];number(original.get('qty'),'原数量');number(original.get('amount'),'原金额')
                    if not original['name']:raise ValueError('原品名缺失，不能生成开票成品')
                    conflict=any(x['conflict'] for x in self.candidates(c,original,ctx))
                else:
                    if c.execute("SELECT 1 FROM imports WHERE status='running' LIMIT 1").fetchone():raise ValueError('历史库仍在导入，请完成后复核')
                validate_yes(r,conflict);status='yes'
                if b and not b['is_test']:warning=self.remember(c,obj,r,b)
            if kind=='group':
                c.execute('UPDATE groups SET review=?,status=?,reason=?,version=version+1,updated=? WHERE id=?',(js(r),status,reason,stamp(),target))
            else:c.execute('UPDATE items SET review=?,status=?,reason=?,version=version+1 WHERE id=?',(js(r),status,reason,target))
            self.lib.audit(c,kind,target,self.actor,action,before,{'review':r,'status':status,'reason':reason,'version':obj['version']+1})
            result=self.lib.group(target,c) if kind=='group' else self.item(c,target)
            result['knowledge_warning']=warning;return result
    def annotate(self,p,c=None):
        if c is None:
            with self.lib.connect() as conn:
                conn.execute('BEGIN IMMEDIATE')
                return self.annotate(p,conn)
        """A lightweight rate label, not a verified invoice or universal tax rule."""
        action=p.get('action')
        if action not in ('yes','no','edit'):raise ValueError('标注操作无效')
        note=text(p.get('note'));chosen=rate(p.get('rate'))
        if len(note)>4000:raise ValueError('备注过长')
        if action=='yes' and chosen not in ('0','0.09','0.13'):raise ValueError('请选择0、9%或13%')
        if action=='no' and not note:raise ValueError('请写一句不通过的原因')
        if action=='yes' and p.get('seen') is not True:raise ValueError('请逐条人工确认')
        obj=self.lib.group(p['id'],c);assert_version(obj,p)
        if c.execute("SELECT 1 FROM imports WHERE status='running' LIMIT 1").fetchone():raise ValueError('历史库仍在导入，请稍后标注')
        before={'review':obj['review'],'status':obj['status'],'reason':obj['reason'],'version':obj['version']}
        # Preserve original facts and detailed prior review, but never manufacture missing facts.
        r={**obj['review'],'name':obj['review'].get('name') or obj['name'],
           'spec':obj['review'].get('spec',obj['spec']),'unit':obj['review'].get('unit',obj['unit']),
           'rate':chosen,'annotation_note':note,'annotation_only':True}
        if 'tax_code' in p:r['tax_code']=text(p['tax_code'])
        if chosen!='0':
            for key in ('treatment','policy','export_policy'):r[key]=''
        status='pending' if action=='edit' else action
        reason=note if action=='no' else ''
        c.execute('UPDATE groups SET review=?,status=?,reason=?,version=version+1,updated=? WHERE id=?',(js(r),status,reason,stamp(),obj['id']))
        self.lib.audit(c,'group',obj['id'],self.actor,'税率标注:'+action,before,{'review':r,'status':status,'reason':reason,'version':obj['version']+1})
        return self.lib.group(obj['id'],c)
    def annotate_bulk(self,p):
        targets=p.get('targets')
        if not isinstance(targets,list) or not 1<=len(targets)<=100:raise ValueError('请明确勾选1至100条商品')
        ids=[t.get('id') for t in targets if isinstance(t,dict)]
        if len(ids)!=len(targets) or len(set(ids))!=len(ids):raise ValueError('选择包含重复或无效商品')
        if p.get('action') not in ('yes','no'):raise ValueError('批量操作无效')
        if p.get('seen') is not True:raise ValueError('请确认已勾选的商品范围')
        with self.lib.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            # All-or-nothing version checks; no wildcard or search-wide updates.
            for t in targets:assert_version(self.lib.group(t['id'],c),t)
            rows=[self.annotate({**p,'id':t['id'],'version':t['version']},c) for t in targets]
        return {'count':len(rows),'rows':rows}
    def remember(self,c,item,r,b):
        ctx=b['context'];o=item['original']
        knowledge={**r,'seller':ctx['seller']}
        identity={k:knowledge.get(k,'') for k in ('identity','facts','process','taxpayer','trade','valid_from','valid_to','seller')}
        gid=digest(['confirmed',o['name'],r['spec'],r['unit'],identity]);old=c.execute('SELECT * FROM groups WHERE id=?',(gid,)).fetchone()
        if old:
            prev=json.loads(old['review']);different=any(prev.get(k)!=knowledge.get(k) for k in ('name','tax_code','rate','treatment','policy','export_policy'))
            if different:
                c.execute('INSERT INTO proposals VALUES(?,?,?,?,?)',(uid(),gid,js(knowledge),self.actor,stamp()))
                c.execute("UPDATE groups SET conflict=1,status='pending',version=version+1 WHERE id=?",(gid,))
                self.lib.audit(c,'group',gid,self.actor,'冲突建议，未覆盖已确认知识',prev,knowledge)
                return '相同商品与适用条件已有不同意见：已保留冲突建议，未覆盖原知识。'
            return '同品同条件知识已存在；本行确认单独保留。'
        c.execute("INSERT INTO groups(id,name,spec,unit,seller,origin,status,review,updated) VALUES(?,?,?,?,?,'confirmed','yes',?,?)",(gid,o['name'],r['spec'],r['unit'],ctx['seller'],js(knowledge),stamp()))
        self.lib.audit(c,'group',gid,self.actor,'业务行人工通过后记入知识',{},knowledge)
        return '已写入同品、同规格、同加工与交易条件的确认知识。'
    def set_template(self,p):
        with self.lib.connect() as c:u=self.uploaded(c,p['token'])
        if not u['name'].endswith('.xlsx'):raise ValueError('模板需xlsx')
        meta=inspect_template(u['path'])
        # Template is account-scoped; histories and reviewed knowledge are shared.
        with self.lib.connect() as c:c.execute('INSERT OR REPLACE INTO settings VALUES(?,?)',('template:'+self.owner,js({'path':u['path'],'sha':u['sha'],'name':u['name']})))
        return {'name':u['name'],'sha':u['sha'],'meta':meta}
    def export(self,p):
        draft=p.get('mode')=='draft'
        if p.get('mode') not in ('draft','final'):raise ValueError('请选择核对稿或最终模板')
        with self.lib.connect() as c:
            c.execute('BEGIN IMMEDIATE');b=self.batch(c,p['id']);assert_version(b,p)
            items=[self.item(c,x[0]) for x in c.execute('SELECT id FROM items WHERE batch_id=? ORDER BY ordinal',(b['id'],))]
            token=uid();out=self.files/(token+('.csv' if draft else '.xlsx'))
            if draft:
                with out.open('w',encoding='utf-8-sig',newline='') as f:
                    w=csv.writer(f);w.writerow(['仅供核对，禁止直接开票',b['risk'] or '所有行保留；未人工通过的字段不是开票结论'])
                    w.writerow(['序号','原表行','原品名','规格','单位','数量','原金额','状态','人工意见','候选名称','候选分类码','候选税率'])
                    for i in items:
                        o=i['original'];r=i['review'];vals=[i['ordinal'],i['source_row'],o['name'],o['spec'],o['unit'],o['qty'],o['amount'],i['status'],i['reason'],r.get('name',''),r.get('tax_code',''),r.get('rate','')]
                        w.writerow(["'"+v if isinstance(v,str) and v[:1] in ('=','+','-','@','\t','\r') else v for v in vals])
                return {'url':'/download/v2/'+out.name,'name':'仅供核对-禁止直接开票.csv','rows':len(items),'draft':True}
            if b['risk']:raise ValueError(b['risk'])
            if not items or any(i['status']!='yes' for i in items):raise ValueError('所有业务行必须新鲜人工通过；未看、No、修改待审均阻断，不能漏行导出')
            ctx=b['context']
            if not ctx.get('basis_confirmed') or ctx.get('amount_basis') not in ('含税','不含税') or not ctx.get('basis_evidence'):raise ValueError('须确认金额含税口径及依据；不自行换算')
            config=c.execute('SELECT value FROM settings WHERE key=?',('template:'+self.owner,)).fetchone()
            if config:tpl=json.loads(config[0])
            else:
                shared=self.lib.path.parent/'tax-template.xlsx'
                if not shared.is_file():raise ValueError('先上传并核验原电子税务局模板')
                tpl={'path':str(shared),'sha':file_hash(shared)}
            if file_hash(tpl['path'])!=tpl['sha']:raise ValueError('模板文件已变化，请重新核验')
            values=[]
            for i in items:
                r=i['review'];validate_yes(r)
                values.append({**r,'qty':i['original']['qty'],'unit_price':'','amount':i['original']['amount']})
            result=write_template(tpl['path'],values,out)
            self.lib.audit(c,'batch',b['id'],self.actor,'导出全部通过行；未验证税局导入',{}, {'rows':len(items),'amount_basis':ctx['amount_basis'],'template_sha':tpl['sha']})
            return {**result,'url':'/download/v2/'+out.name,'name':'税务模板-需税局校验.xlsx','amount_basis':ctx['amount_basis'],'draft':False}
    def agent_export(self,p):
        kind=p.get('kind','group');offset=max(0,int(p.get('offset',0)));limit=max(1,min(int(p.get('limit',500)),500))
        with self.lib.connect() as c:
            if kind=='group':ids=[r[0] for r in c.execute("SELECT id FROM groups WHERE status='no' ORDER BY id LIMIT ? OFFSET ?",(limit,offset))]
            elif kind=='item':
                self.batch(c,p['batch_id']);ids=[r[0] for r in c.execute("SELECT id FROM items WHERE batch_id=? AND status='no' ORDER BY ordinal LIMIT ? OFFSET ?",(p['batch_id'],limit,offset))]
            else:raise ValueError('对象无效')
            entries=[];targets={}
            for iid in ids:
                obj=self.lib.group(iid,c) if kind=='group' else self.item(c,iid)
                original=obj if kind=='group' else obj['original'];history=obj.get('history',[])
                if kind=='item':history=[h for g in self.candidates(c,original,{}) for h in g['history']]
                reason=re.sub(r'\b[A-Za-z0-9]{15,}\b','[标识已移除]',obj['reason'])
                # No source file/path, counterparty, invoice IDs, quantities, amounts or credentials.
                for s in c.execute("SELECT DISTINCT seller FROM groups WHERE seller<>''"):
                    reason=reason.replace(s[0],'[主体已移除]')
                entries.append({'id':iid,'version':obj['version'],'name':original['name'],'spec':original['spec'],'unit':original['unit'],'historical_candidates':[{'tax_code':h['code'],'rate':h['rate']} for h in history],'human_reason':reason})
                targets[iid]={'kind':kind,'version':obj['version']}
            if not entries:raise ValueError('当前没有No待核记录')
            bid=uid();c.execute('INSERT INTO bundles VALUES(?,?,?,?)',(bid,self.owner,js(targets),stamp()))
        payload={'schema':'gc-tax-agent-v1','bundle_id':bid,'instructions':'仅建议同一真实商品，先辨商品再给编码税率；不替换交易品目。回填id/version/name/tax_code/rate/reason/basis；不会自动通过。不要添加客户身份和账务。','entries':entries,'next_offset':offset+len(entries)}
        out=self.files/(bid+'.json');out.write_text(js(payload),encoding='utf-8')
        return {'url':'/download/v2/'+out.name,'name':'Agent待核建议包.json','count':len(entries),'next_offset':payload['next_offset']}
    def agent_import(self,p):
        payload=p['payload']
        if payload.get('schema')!='gc-tax-agent-v1' or not isinstance(payload.get('entries'),list):raise ValueError('不是约定的Agent建议包')
        if not 1<=len(payload['entries'])<=500:raise ValueError('每包1至500条')
        with self.lib.connect() as c:
            c.execute('BEGIN IMMEDIATE');bundle=c.execute('SELECT * FROM bundles WHERE id=? AND owner=?',(payload.get('bundle_id'),self.owner)).fetchone()
            if not bundle:raise ValueError('此账号无对应导出包')
            targets=json.loads(bundle['targets']);seen=set()
            for entry in payload['entries']:
                iid=entry.get('id');target=targets.get(iid)
                if not target or iid in seen:raise ValueError('建议包含重复或非原包ID')
                seen.add(iid);kind=target['kind'];obj=self.lib.group(iid,c) if kind=='group' else self.item(c,iid)
                if entry.get('version')!=target['version'] or obj['version']!=target['version']:raise ValueError('Agent建议版本过期，整包未覆盖；请重新导出')
                for k in ('name','tax_code','rate','reason','basis'):
                    if not text(entry.get(k)):raise ValueError('Agent建议缺少'+k)
                r={**obj['review'],'name':text(entry['name']),'tax_code':text(entry['tax_code']),'rate':rate(entry['rate']),'basis':text(entry['basis']),'agent_reason':text(entry['reason']),'same_product':False,'conditions_checked':False}
                table='groups' if kind=='group' else 'items'
                c.execute('UPDATE '+table+" SET review=?,status='pending',version=version+1 WHERE id=?",(js(r),iid))
                self.lib.audit(c,kind,iid,self.actor,'导入Agent建议，仍待人工确认',obj['review'],r)
        return {'imported':len(seen),'auto_confirmed':0}
    def get(self,path,p):
        if path=='stats':return self.lib.stats()
        if path=='groups':return self.lib.search(p.get('q',''),p.get('status',''),p.get('offset',0),p.get('limit',30),**candidate_options(p))
        if path=='group':return self.lib.group(p['id'])
        if path=='sources':return self.lib.sources(p['id'],p.get('offset',0),p.get('limit',30))
        if path=='batch':return self.batch_view(p)
        if path=='upload-info':return self.upload_info(p)
        with self.lib.connect() as c:
            if path=='batches':return {'rows':[dict(r) for r in c.execute('SELECT b.id,b.title,b.version,b.is_test,b.risk,b.created,COUNT(i.id) AS total,SUM(i.status=\'yes\') AS yes FROM batches b LEFT JOIN items i ON i.batch_id=b.id WHERE owner=? GROUP BY b.id ORDER BY b.created DESC',(self.owner,))]}
            if path=='audit':
                kind=p['kind']
                if kind=='item':self.item(c,p['id'])
                elif kind=='batch':self.batch(c,p['id'])
                elif kind!='group':raise ValueError('对象无效')
                return {'rows':[{**dict(r),'before':json.loads(r['before_json']),'after':json.loads(r['after_json'])} for r in c.execute('SELECT * FROM audit WHERE kind=? AND target=? ORDER BY seq DESC LIMIT 100',(kind,p['id']))]}
            if path=='template':
                row=c.execute('SELECT value FROM settings WHERE key=?',('template:'+self.owner,)).fetchone()
                shared=self.lib.path.parent/'tax-template.xlsx'
                return {'configured':bool(row) or shared.is_file(),'name':json.loads(row[0])['name'] if row else ('159962.6.xlsx · 已核共享模板' if shared.is_file() else '')}
        raise ValueError('接口不存在')
    def post(self,path,p):
        methods={'annotate-bulk':self.annotate_bulk,'annotate':self.annotate,'upload-start':self.upload_start,'upload-chunk':self.upload_chunk,'upload-info':self.upload_info,'history-import':self.start_history,'batch-import':self.import_batch,'context':self.save_context,'review':self.review,'template':self.set_template,'export':self.export,'agent-export':self.agent_export,'agent-import':self.agent_import}
        if path not in methods:raise ValueError('接口不存在')
        return methods[path](p)
