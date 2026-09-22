"""Private, streaming history library. No model, ERP or invoice API calls."""
import datetime as dt
import hashlib
import json
import re
import sqlite3
from contextlib import contextmanager
from pathlib import Path
import openpyxl

KNOWN = {
    'e9ad92af258c1a999b0427ee1b0704fe513c8150b010c68edf3578393d380136': {'Sheet1':129382},
    '893d23b09c5039bb006791321da0f46ea839dbb1e2a96130c5aefadb432e0d99': {'信息汇总表1':102614},
}
CONSISTENT_SQL = """(conflict=0 AND EXISTS(SELECT 1 FROM observations o WHERE o.group_id=groups.id)
AND NOT EXISTS(SELECT 1 FROM observations o WHERE o.group_id=groups.id AND
(o.rate NOT IN ('0','0.09','0.13') OR length(o.code)!=19 OR o.code GLOB '*[^0-9]*')))"""
SCHEMA = '''
CREATE TABLE IF NOT EXISTS imports(sha TEXT PRIMARY KEY,name TEXT,kind TEXT,status TEXT,counts TEXT DEFAULT '{}',error TEXT DEFAULT '',created TEXT);
CREATE TABLE IF NOT EXISTS groups(id TEXT PRIMARY KEY,name TEXT,spec TEXT,unit TEXT,seller TEXT DEFAULT '',origin TEXT DEFAULT 'history',status TEXT DEFAULT 'pending',version INTEGER DEFAULT 1,review TEXT DEFAULT '{}',reason TEXT DEFAULT '',conflict INTEGER DEFAULT 0,updated TEXT);
CREATE INDEX IF NOT EXISTS group_name ON groups(name,spec,unit);
CREATE INDEX IF NOT EXISTS group_review_name ON groups(json_extract(review,'$.name')) WHERE status='yes';
CREATE INDEX IF NOT EXISTS group_status ON groups(status);
CREATE TABLE IF NOT EXISTS observations(id TEXT PRIMARY KEY,group_id TEXT REFERENCES groups(id),code TEXT,rate TEXT,raw TEXT);
CREATE INDEX IF NOT EXISTS obs_group ON observations(group_id);
CREATE INDEX IF NOT EXISTS obs_group_values ON observations(group_id,code,rate);
CREATE INDEX IF NOT EXISTS group_review_state ON groups(status,conflict,id);
CREATE INDEX IF NOT EXISTS group_browse ON groups(name,spec,unit,id,status,conflict);
CREATE TABLE IF NOT EXISTS sources(sha TEXT REFERENCES imports(sha),sheet TEXT,row_no INTEGER,group_id TEXT REFERENCES groups(id),obs_id TEXT REFERENCES observations(id),kind TEXT,alias TEXT,full_name TEXT,PRIMARY KEY(sha,sheet,row_no));
CREATE INDEX IF NOT EXISTS source_group ON sources(group_id,sha,sheet,row_no);
CREATE TABLE IF NOT EXISTS observation_dates(id TEXT PRIMARY KEY REFERENCES observations(id),source_date TEXT);
CREATE INDEX IF NOT EXISTS observation_date ON observation_dates(source_date,id);
CREATE TABLE IF NOT EXISTS audit(seq INTEGER PRIMARY KEY,kind TEXT,target TEXT,actor TEXT,action TEXT,before_json TEXT,after_json TEXT,created TEXT);
CREATE INDEX IF NOT EXISTS audit_target ON audit(kind,target,seq);
CREATE TABLE IF NOT EXISTS proposals(id TEXT PRIMARY KEY,group_id TEXT,review TEXT,actor TEXT,created TEXT);
CREATE TABLE IF NOT EXISTS batches(id TEXT PRIMARY KEY,owner TEXT,title TEXT,version INTEGER DEFAULT 1,context TEXT DEFAULT '{}',risk TEXT DEFAULT '',is_test INTEGER DEFAULT 0,source_sha TEXT,created TEXT);
CREATE INDEX IF NOT EXISTS batch_owner ON batches(owner);
CREATE TABLE IF NOT EXISTS items(id TEXT PRIMARY KEY,batch_id TEXT REFERENCES batches(id),ordinal INTEGER,source_row INTEGER,original TEXT,status TEXT DEFAULT 'pending',version INTEGER DEFAULT 1,review TEXT DEFAULT '{}',reason TEXT DEFAULT '');
CREATE INDEX IF NOT EXISTS item_batch ON items(batch_id,ordinal);
CREATE TABLE IF NOT EXISTS bundles(id TEXT PRIMARY KEY,owner TEXT,targets TEXT,created TEXT);
CREATE TABLE IF NOT EXISTS uploads(id TEXT PRIMARY KEY,owner TEXT,name TEXT,path TEXT,size INTEGER DEFAULT 0,sha TEXT DEFAULT '',complete INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT);
'''

def stamp():return dt.datetime.now().astimezone().isoformat(timespec='seconds')
def js(v):return json.dumps(v,ensure_ascii=False,sort_keys=True,separators=(',',':'),default=str)
def digest(v):return hashlib.sha256(js(v).encode()).hexdigest()
def file_hash(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()
def text(v):
    if v is None:return ''
    if isinstance(v,(dt.datetime,dt.date)):return v.isoformat()
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return str(v).strip()
def rate(v):
    s=text(v)
    if not s:return ''
    try:
        n=float(s.rstrip('%'));n=n/100 if '%' in s or n>1 else n
        return format(n,'.8g')
    except ValueError:return s
def strip_category(value):
    s=text(value)
    return re.sub(r'^\*[^*]+\*','',s).strip()

def source_date(value):
    # Only the explicit invoice-date field; never infer from file mtime/import time.
    value=text(value)
    match=re.fullmatch(r'(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})(?:日)?(?:[ T]\d{2}:\d{2}:\d{2}(?:\.\d+)?)?',value)
    if not match:return None
    try:return dt.date(*map(int,match.groups())).isoformat()
    except ValueError:return None

def candidate_options(p):
    out={k:text(p.get(k,'')) for k in ('date_from','date_to','date_mode','sort')}
    for key in ('date_from','date_to'):
        if out[key] and (source_date(out[key])!=out[key]):raise ValueError('来源日期须为有效 YYYY-MM-DD')
    if out['date_from'] and out['date_to'] and out['date_from']>out['date_to']:raise ValueError('来源起日不能晚于止日')
    if out['date_mode'] not in ('','known','unknown'):raise ValueError('来源日期范围无效')
    if out['sort'] not in ('','recent','frequency'):raise ValueError('候选排序无效')
    return out

def date_clause(options,alias='d'):
    parts=[];args=[]
    for key,op in (('date_from','>='),('date_to','<=')):
        if options.get(key):parts.append(f'{alias}.source_date{op}?');args.append(options[key])
    if options.get('date_mode')=='unknown':parts.append(f'{alias}.source_date IS NULL')
    elif options.get('date_mode')=='known':parts.append(f'{alias}.source_date IS NOT NULL')
    return ' AND '.join(parts) or '1=1',args

class Library:
    def __init__(self,path):
        self.path=Path(path);self.path.parent.mkdir(parents=True,exist_ok=True)
        with self.connect() as c:
            c.executescript(SCHEMA)
            if not c.execute("SELECT 1 FROM settings WHERE key='source-date-v1'").fetchone():
                rows=c.execute('SELECT id,raw FROM observations').fetchall()
                c.executemany('INSERT OR IGNORE INTO observation_dates VALUES(?,?)',[(r['id'],source_date(json.loads(r['raw']).get('date'))) for r in rows])
                c.execute("INSERT INTO settings VALUES('source-date-v1','complete')")
    @contextmanager
    def connect(self):
        c=sqlite3.connect(self.path,timeout=30);c.row_factory=sqlite3.Row
        c.execute('PRAGMA journal_mode=WAL');c.execute('PRAGMA foreign_keys=ON')
        try:
            yield c;c.commit()
        except BaseException:c.rollback();raise
        finally:c.close()
    def audit(self,c,kind,target,actor,action,before,after):
        c.execute('INSERT INTO audit(kind,target,actor,action,before_json,after_json,created) VALUES(?,?,?,?,?,?,?)',(kind,target,actor,action,js(before),js(after),stamp()))
    def add_observation(self,c,sha,sheet,n,fields,raw_row,kind='detail',paired=None):
        identity={k:fields[k] for k in ['name','spec','unit','seller','business']}
        gid=digest(['history',identity]);oid=digest(['source-row',raw_row])
        if paired:gid,oid=paired
        else:
            c.execute('INSERT OR IGNORE INTO groups(id,name,spec,unit,seller,updated) VALUES(?,?,?,?,?,?)',(gid,fields['name'],fields['spec'],fields['unit'],fields['seller'],stamp()))
            c.execute('INSERT OR IGNORE INTO observations VALUES(?,?,?,?,?)',(oid,gid,fields['code'],fields['rate'],js(fields)))
        c.execute('INSERT OR IGNORE INTO observation_dates VALUES(?,?)',(oid,source_date(fields.get('date'))))
        c.execute('INSERT OR IGNORE INTO sources VALUES(?,?,?,?,?,?,?,?)',(sha,sheet,n,gid,oid,kind,fields['alias'],fields['full_name']))
    def import_history(self,path,expected_sha=None):
        path=Path(path);sha=file_hash(path)
        if expected_sha and sha!=expected_sha:raise ValueError('原件哈希不符，停止导入')
        with self.connect() as c:
            old=c.execute('SELECT * FROM imports WHERE sha=?',(sha,)).fetchone()
            if old and old['status']=='complete':return {**dict(old),'counts':json.loads(old['counts']),'idempotent':True}
            c.execute("INSERT INTO imports(sha,name,kind,status,created) VALUES(?,?,'history','running',?) ON CONFLICT(sha) DO UPDATE SET status='running',error=''",(sha,path.name,stamp()))
        counts={};book=None
        try:
            book=openpyxl.load_workbook(path,read_only=True,data_only=True)
            if '信息汇总表1' not in book.sheetnames:raise ValueError('历史库需包含 信息汇总表1；请勿将业务开票清单当作历史库')
            sheets=['信息汇总表1']+(['Sheet1'] if 'Sheet1' in book.sheetnames else [])
            for name in sheets:
                s=book[name];it=s.iter_rows(values_only=True);headers=[text(v) for v in next(it)]
                def col(names,default=-1):return next((headers.index(k) for k in names if k in headers),default)
                alias_i=col(['商品名称','名称']);full_i=col(['货物或应税劳务名称'])
                code_i=col(['税收分类编码']);rate_i=col(['税率']);spec_i=col(['规格型号','规格'],14 if sha in KNOWN and name=='信息汇总表1' and len(headers)==29 else -1)
                unit_i=col(['单位']);seller_i=col(['销方名称']);date_i=col(['开票日期']);business_i=col(['特定业务类型'])
                if alias_i<0 or code_i<0 or rate_i<0:raise ValueError('缺少名称、税收分类编码或税率表头')
                counts[name]={'named_rows':0,'source_refs':0,'kind':'reference' if name=='Sheet1' else 'detail','linked_detail_refs':0}
                # Commit bounded batches; no worksheet-wide list or row limit.
                with self.connect() as c:
                    for n,row in enumerate(it,2):
                        get=lambda i:text(row[i]) if 0<=i<len(row) else ''
                        alias=get(alias_i);full=get(full_i)
                        if not alias:continue
                        counts[name]['named_rows']+=1
                        full_clean=strip_category(full)
                        canonical=full_clean if full.startswith('*') and re.match(r'^\*[^*]+\*.+',full) else alias
                        fields={'name':canonical,'alias':alias,'full_name':full,'spec':get(spec_i),'unit':get(unit_i),'seller':get(seller_i),'date':get(date_i),'business':get(business_i),'code':get(code_i),'rate':rate(get(rate_i)),'rate_original':get(rate_i)}
                        paired=None
                        if name=='Sheet1':
                            # Link ONLY a verified corresponding source row, never by an ambiguous shortened name.
                            prior=c.execute("SELECT s.group_id,s.obs_id,s.alias,o.code,o.rate FROM sources s JOIN observations o ON o.id=s.obs_id WHERE s.sha=? AND s.sheet='信息汇总表1' AND s.row_no=?",(sha,n)).fetchone()
                            if prior and (prior['alias'],prior['code'],prior['rate'])==(alias,fields['code'],fields['rate']):
                                paired=(prior['group_id'],prior['obs_id']);counts[name]['linked_detail_refs']+=1
                        self.add_observation(c,sha,name,n,fields,list(row),'reference' if name=='Sheet1' else 'detail',paired)
                        counts[name]['source_refs']+=1
                        if counts[name]['named_rows']%1000==0:
                            c.execute('UPDATE imports SET counts=? WHERE sha=?',(js(counts),sha));c.commit()
                    c.execute('UPDATE imports SET counts=? WHERE sha=?',(js(counts),sha))
            for sheet,expected in KNOWN.get(sha,{}).items():
                if counts.get(sheet,{}).get('named_rows')!=expected:raise ValueError('原件含名称行数与基线不符：'+sheet)
            with self.connect() as c:
                # More than one historical code/rate is a conflict, never majority voting.
                c.execute("UPDATE groups SET conflict=(SELECT CASE WHEN COUNT(DISTINCT code || '|' || rate)>1 THEN 1 ELSE 0 END FROM observations WHERE group_id=groups.id) WHERE origin='history'")
                c.execute("UPDATE groups SET status='pending',version=version+1,reason='新增历史来源，请重新核实',updated=? WHERE status='yes' AND origin='history' AND id IN (SELECT group_id FROM sources WHERE sha=?)",(stamp(),sha))
                c.execute("UPDATE imports SET status='complete',counts=? WHERE sha=?",(js(counts),sha))
            return {'sha':sha,'name':path.name,'counts':counts,'status':'complete','idempotent':False}
        except Exception as e:
            with self.connect() as c:c.execute("UPDATE imports SET status='error',error=? WHERE sha=?",(str(e),sha))
            raise
        finally:
            if book:book.close()

    def stats(self):
        with self.connect() as c:
            states={r[0]:r[1] for r in c.execute('SELECT status,COUNT(*) FROM groups GROUP BY status')}
            consistent=c.execute("SELECT COUNT(*) FROM groups WHERE status='pending' AND "+CONSISTENT_SQL).fetchone()[0]
            diffs=c.execute("SELECT SUM(rates>1),SUM(rates=1 AND codes>1) FROM (SELECT COUNT(DISTINCT o.rate) rates,COUNT(DISTINCT o.code) codes FROM groups g JOIN observations o ON o.group_id=g.id WHERE g.conflict=1 GROUP BY g.id)").fetchone()
            return {'rate_conflict':diffs[0] or 0,'code_only':diffs[1] or 0,'total':sum(states.values()),'pending':states.get('pending',0),'yes':states.get('yes',0),'no':states.get('no',0),'conflict':c.execute('SELECT COUNT(*) FROM groups WHERE conflict=1').fetchone()[0],
                    'consistent':consistent,
                    'needs_review':states.get('pending',0)-consistent+states.get('no',0),
                    'source_refs':c.execute('SELECT COUNT(*) FROM sources').fetchone()[0],'observations':c.execute('SELECT COUNT(*) FROM observations').fetchone()[0],
                    'imports':[{**dict(r),'counts':json.loads(r['counts'])} for r in c.execute('SELECT * FROM imports ORDER BY created')],'model':'未接模型；历史记录不是税务结论，汇总引用不重复计交易'}

    def group(self,gid,c=None):
        if c is None:
            with self.connect() as conn:return self.group(gid,conn)
        r=c.execute('SELECT * FROM groups WHERE id=?',(gid,)).fetchone()
        if not r:raise ValueError('商品组不存在')
        r=dict(r);r['review']=json.loads(r['review'])
        r['history']=[dict(x) for x in c.execute('SELECT code,rate,COUNT(*) AS observations FROM observations WHERE group_id=? GROUP BY code,rate ORDER BY code,rate',(gid,))]
        metrics=c.execute('SELECT COUNT(*) occurrences,MAX(d.source_date) last_used,SUM(d.source_date IS NULL) unknown_dates FROM observations o LEFT JOIN observation_dates d ON d.id=o.id WHERE o.group_id=?',(gid,)).fetchone()
        r.update(dict(metrics))
        r['date_basis']='原表开票日期；未知日期不以导入时间补齐'
        r['count_basis']='去重历史记录数，汇总引用不重复计数；不等于独立交易次数'
        r['source_files']=[dict(x) for x in c.execute('SELECT DISTINCT i.name file,s.sheet FROM sources s JOIN imports i ON i.sha=s.sha WHERE s.group_id=? ORDER BY i.name,s.sheet',(gid,))]
        r['confirmation']=dict(c.execute("SELECT actor,created,action FROM audit WHERE kind='group' AND target=? ORDER BY seq DESC LIMIT 1",(gid,)).fetchone() or {})
        r['source_count']=c.execute('SELECT COUNT(*) FROM sources WHERE group_id=?',(gid,)).fetchone()[0]
        r['proposals']=[{**dict(x),'review':json.loads(x['review'])} for x in c.execute('SELECT * FROM proposals WHERE group_id=? ORDER BY created DESC LIMIT 100',(gid,))]
        flags=[]
        if not r['spec']:flags.append('规格缺失，需明确商品事实')
        if not r['unit']:flags.append('单位缺失')
        if r['conflict']:flags.append('编码/税率或人工意见有差异，禁止多数投票')
        if any(x['rate']=='0' for x in r['history']):flags.append('历史0含义待核')
        if any(x['rate'] not in ('0','0.09','0.13','') for x in r['history']):flags.append('历史税率超出0/9%/13%，保留原值')
        if c.execute("SELECT 1 FROM sources WHERE group_id=? AND full_name<>'' AND alias<>'' AND full_name NOT LIKE '%'||alias||'%' LIMIT 1",(gid,)).fetchone():flags.append('原名称与完整货物名不一致')
        r['rate_conflict']=len({x['rate'] for x in r['history']})>1
        r['code_conflict']=len({x['code'] for x in r['history']})>1
        r['history_consistent']=not r['conflict'] and bool(r['history']) and all(x['rate'] in ('0','0.09','0.13') and re.fullmatch(r'\d{19}',x['code'] or '') for x in r['history'])
        r['flags']=flags
        return r
    def search(self,q='',status='',offset=0,limit=30,**options):
        options=candidate_options(options);dc,da=date_clause(options)
        limit=max(1,min(int(limit),100));offset=max(0,int(offset));where=['1=1'];args=[]
        for term in text(q).split():
            where.append('(name LIKE ? OR spec LIKE ? OR id=?)');args.extend(['%'+term+'%','%'+term+'%',term])
        if status=='needs_review':where.append("(status='no' OR (status='pending' AND NOT "+CONSISTENT_SQL+"))")
        elif status=='consistent':where.append("status='pending' AND "+CONSISTENT_SQL)
        elif status=='rate_conflict':where.append("conflict=1 AND (SELECT COUNT(DISTINCT rate) FROM observations o WHERE o.group_id=groups.id)>1")
        elif status=='code_only':where.append("conflict=1 AND (SELECT COUNT(DISTINCT rate) FROM observations o WHERE o.group_id=groups.id)=1 AND (SELECT COUNT(DISTINCT code) FROM observations o WHERE o.group_id=groups.id)>1")
        elif status=='conflict':where.append('conflict=1')
        elif status in ('pending','yes','no'):where.append('status=?');args.append(status)
        if dc!='1=1':
            where.append('EXISTS(SELECT 1 FROM observations o LEFT JOIN observation_dates d ON d.id=o.id WHERE o.group_id=groups.id AND '+dc+')');args.extend(da)
        clause=' AND '.join(where)
        scoped="SELECT o.group_id,COUNT(*) occurrences,MAX(d.source_date) last_used FROM observations o LEFT JOIN observation_dates d ON d.id=o.id WHERE "+dc+" GROUP BY o.group_id"
        ordering={'recent':'m.last_used DESC,','frequency':'m.occurrences DESC,'}.get(options['sort'],'')+'name,spec,unit,id'
        with self.connect() as c:
            total=c.execute('SELECT COUNT(*) FROM groups WHERE '+clause,args).fetchone()[0]
            if options['sort']:
                ids=[r[0] for r in c.execute('WITH m AS ('+scoped+') SELECT groups.id FROM groups LEFT JOIN m ON m.group_id=groups.id WHERE '+clause+' ORDER BY '+ordering+' LIMIT ? OFFSET ?',da+args+[limit,offset])]
            else:
                ids=[r[0] for r in c.execute('SELECT id FROM groups WHERE '+clause+' ORDER BY name,spec,unit,id LIMIT ? OFFSET ?',args+[limit,offset])]
            rows=[self.group(gid,c) for gid in ids]
            for row in rows:
                if dc!='1=1':
                    row['all_occurrences']=row['occurrences']
                    row.update(dict(c.execute('SELECT COUNT(*) occurrences,MAX(d.source_date) last_used FROM observations o LEFT JOIN observation_dates d ON d.id=o.id WHERE o.group_id=? AND '+dc,[row['id']]+da).fetchone()))
            return {'total':total,'offset':offset,'limit':limit,'rows':rows}
    def sources(self,gid,offset=0,limit=30):
        limit=max(1,min(int(limit),100));offset=max(0,int(offset))
        with self.connect() as c:
            return {'total':c.execute('SELECT COUNT(*) FROM sources WHERE group_id=?',(gid,)).fetchone()[0],'offset':offset,
              'files':[dict(r) for r in c.execute('SELECT i.name AS file,s.sheet,COUNT(*) AS rows FROM sources s JOIN imports i ON i.sha=s.sha WHERE s.group_id=? GROUP BY s.sha,s.sheet ORDER BY i.name,s.sheet',(gid,))],
              'rows':[{**dict(r),'raw':json.loads(r['raw'])} for r in c.execute('SELECT s.sha,i.name AS file,s.sheet,s.row_no,s.kind,s.alias,s.full_name,o.raw FROM sources s JOIN imports i ON i.sha=s.sha JOIN observations o ON o.id=s.obs_id WHERE s.group_id=? ORDER BY s.sha,s.sheet,s.row_no LIMIT ? OFFSET ?',(gid,limit,offset))]}
