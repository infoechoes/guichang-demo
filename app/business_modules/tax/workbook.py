"""Write review workbooks with the Excel dependency already used by the server."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def write_workbook(data, output):
    book = Workbook()
    sheet = book.active
    sheet.title = data.get('sheet', '税率复核')

    def put(ws, row, values):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            # Preserve imported strings, including leading zeroes and '=' text.
            if isinstance(value, str):
                cell.data_type = 's'
            cell.font = Font(name='Microsoft YaHei', size=10, color='26372F')
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    put(sheet, 1, [data['title']])
    put(sheet, 2, [data['note']])
    put(sheet, 4, data['headers'])
    sheet['A1'].font = Font(name='Microsoft YaHei', size=14, bold=True)
    sheet['A2'].font = Font(name='Microsoft YaHei', size=10, color='737972')
    for row, values in enumerate(data['rows'], 5):
        put(sheet, row, values)
        sheet.row_dimensions[row].height = 90
        for cell in sheet[row]:
            cell.fill = PatternFill('solid', fgColor='F4F6F0' if row % 2 else 'FFFFFF')
        for col in data.get('percent_columns', []):
            sheet.cell(row, col + 1).number_format = '0%'
    for col in range(len(data['headers'])):
        width = 25 if col in [2, 5, 10, 11, 23, 24, 26, 27] else 16
        if col in [12, 16, 17, 18, 19, 20, 21, 22, 25, 28, 30]:
            width = 48
        sheet.column_dimensions[get_column_letter(col + 1)].width = width
    sheet.column_dimensions['W'].width = 110
    sheet.freeze_panes = 'A5'
    sheet.sheet_view.showGridLines = False
    for cell in sheet[4]:
        cell.fill = PatternFill('solid', fgColor='183B32')
        cell.font = Font(name='Microsoft YaHei', size=10, bold=True, color='FFFFFF')
    sheet.row_dimensions[4].height = 30
    table = Table(displayName='ReviewRecords', ref=f'A4:{get_column_letter(len(data["headers"]))}{len(data["rows"]) + 4}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    sheet.add_table(table)
    if data.get('provenance'):
        trace = book.create_sheet('来源行')
        put(trace, 1, ['样本或记录编号', '商品编码', '商品名称', '来源文件', '来源工作表', '原表行号'])
        for row, values in enumerate(data['provenance'], 2):
            put(trace, row, values)
        for cell in trace[1]:
            cell.fill = PatternFill('solid', fgColor='183B32')
            cell.font = Font(name='Microsoft YaHei', size=10, bold=True, color='FFFFFF')
        for col in range(1, 7):
            trace.column_dimensions[get_column_letter(col)].width = 40 if col == 4 else 26
        trace.freeze_panes = 'A2'
        trace.sheet_view.showGridLines = False
    book.save(output)
    book.close()
