#!/usr/bin/env python3
"""Local-only tax review demo. Never connects to an ERP or external model."""
import base64, copy, csv, datetime as dt, hashlib, io, json, os, re, threading, uuid
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse, parse_qs, quote
import openpyxl
from .policy import SOURCES, VERSION, candidate
from .workbook import write_workbook
from .workflow import Workflow
SHARED_DB=None
OWNER="standalone"
ACTOR="standalone"
def workflow():return Workflow(SHARED_DB or DATA/"tax-knowledge.sqlite3",OWNER,DATA,ACTOR)

ROOT=Path(__file__).resolve().parent
DATA=Path(os.environ.get('TAX_DEMO_DATA', ROOT/'data')).resolve()
PORT=int(os.environ.get('TAX_DEMO_PORT','4296'))
SAMPLE=Path(os.environ['GC_TAX_SAMPLE']) if os.environ.get('GC_TAX_SAMPLE') else None
LOCK=threading.RLock()
RATES=['9%','13%','6%','3%','1%','免税','零税率','不征税','其他']
FIELDS={'code':'商品编码','name':'商品名称','spec':'规格','unit':'单位','rate':'原税率','date':'交易日期','seller':'销售主体','facts':'商品事实','identity':'确认的商品身份'}
ALIASES={'code':['商品编码','商品编号','ERP商品编码','SKU编码'],'name':['商品名称','名称','货物名称'],'spec':['规格','规格型号'],'unit':['单位','计量单位'],'rate':['原税率','税率','销项税率'],'date':['交易日期','日期','开票日期'],'seller':['销售主体','销方名称','销售方名称'],'facts':['商品事实'],'identity':['确认的商品身份']}

def now(): return dt.datetime.now().astimezone().isoformat(timespec='seconds')
def export_time(value):
    if not value: return ''
    parsed=dt.datetime.fromisoformat(value)
    zone=parsed.strftime('%z')
    return parsed.strftime('%Y-%m-%d %H:%M:%S')+'（时区 '+zone[:3]+':'+zone[3:]+'）'
def uid(): return uuid.uuid4().hex
def dump(path,obj):
    path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix('.tmp')
    temp.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8'); temp.replace(path)
def read(path): return json.loads(path.read_text(encoding='utf-8'))
def identified(folder,token,suffix='.json'):
    if not re.fullmatch(r'[a-f0-9]{32}',token or ''): raise ValueError('文件编号无效')
    return DATA/folder/(token+suffix)
def get_batch(token): return read(identified('batches',token))
def persist(b):
    b['revision']+=1; b['saved_at']=now(); dump(identified('batches',b['id']),b)
def fresh(title,rows,context=None):
    return {'schema':1,'id':uid(),'title':title,'revision':0,'created_at':now(),'saved_at':now(),'policy_version':VERSION,'context':context or {'direction':'销项','seller':'','taxpayer':'未核实','trade':'','reviewer':''},'rows':rows,'audit':[]}
def view(b):
    b=copy.deepcopy(b)
    for row in b['rows']: row['candidate']=candidate(row,b['context']['direction'])
    return b
def check_rev(b,p):
    if b['revision']!=p.get('revision'): raise ValueError('此批次已在其他页面更新，请重新打开后再修改，避免覆盖。')
def cell_value(v):
    if isinstance(v,(dt.datetime,dt.date)): return v.isoformat()
    return v
def cell_text(c):
    if c['value'] is None: return ''
    v=c['value']
    if isinstance(v,(int,float)) and re.fullmatch(r'0+',c.get('format','')):
        return str(int(v)).zfill(len(c['format']))
    return str(v)
def parse_file(name,raw):
    if len(raw)>30*1024*1024: raise ValueError('文件超过30MB，请先另存要维护的工作表。')
    if name.lower().endswith('.csv'):
        try: content=raw.decode('utf-8-sig')
        except UnicodeDecodeError: content=raw.decode('gb18030')
        rows=[[{'value':v,'type':'s','format':'General'} for v in r] for r in csv.reader(io.StringIO(content))]
        sheets=[{'name':'CSV','rows':rows,'merged':[]}]
    elif name.lower().endswith('.xlsx'):
        book=openpyxl.load_workbook(io.BytesIO(raw),data_only=False,read_only=False)
        sheets=[]
        for s in book:
            if s.max_row>5000 or s.max_column>200: raise ValueError('旧版样本工具仅用于旧批次；大表请返回新版税率知识库流式导入，不需拆表。')
            rows=[[{'value':cell_value(c.value),'type':c.data_type,'format':c.number_format} for c in r] for r in s.iter_rows()]
            while rows and all(c['value'] is None for c in rows[-1]): rows.pop()
            sheets.append({'name':s.title,'rows':rows,'merged':[str(r) for r in s.merged_cells.ranges]})
        book.close()
    else: raise ValueError('支持 .xlsx 和 .csv；旧版 .xls 请在Excel中另存为 .xlsx。')
    return {'name':Path(name).name,'sha256':hashlib.sha256(raw).hexdigest(),'sheets':sheets}
def upload(p):
    try: raw=base64.b64decode(p['data'],validate=True)
    except Exception: raise ValueError('文件内容无法读取')
    f=parse_file(p['name'],raw); token=uid(); dump(identified('uploads',token),f)
    identified('uploads',token,'.original').write_bytes(raw)
    return {'token':token,'name':f['name'],'sha256':f['sha256'],'sheets':[{'name':s['name'],'rows':len(s['rows']),'cols':max(map(len,s['rows']),default=0),'preview':[[cell_text(c) for c in row] for row in s['rows'][:8]]} for s in f['sheets']]}
def sheet_info(p):
    f=read(identified('uploads',p['token'])); s=next((s for s in f['sheets'] if s['name']==p['sheet']),None)
    if not s: raise ValueError('未找到工作表')
    h=int(p.get('header',1))-1
    if not 0<=h<len(s['rows']): raise ValueError('表头行不在文件范围内')
    headers=[cell_text(c).strip() for c in s['rows'][h]]
    return f,s,h,headers
def preview(p):
    f,s,h,heads=sheet_info(p)
    mapping={k:next((i for i,v in enumerate(heads) if v in names),-1) for k,names in ALIASES.items()}
    if p.get('direction')=='进项': mapping['rate']=next((i for i,v in enumerate(heads) if v=='进项税率'),mapping['rate'])
    return {'headers':heads,'mapping':mapping,'preview':[[cell_text(c) for c in r] for r in s['rows'][h+1:h+6]],'count':len(s['rows'])-h-1}
def rate_values(v):
    if v is None or str(v).strip()=='': return []
    return [str(v)]
def normalized_rate(v):
    s=str(v).strip()
    if s=='免税': return '免税'
    if s in ['零税率','0%']: return '0(含义待核)' if s=='0%' else s
    try:
        n=float(s.rstrip('%')); n=n/100 if '%' in s or n>1 else n
        return '0(含义待核)' if n==0 else f'{n*100:g}%'
    except ValueError: return s
def enrich_flags(rows):
    groups={}; codes={}
    for r in rows:
        # This groups only possible conflicts, not identity or tax correctness.
        key=(r.get('seller',''),r['code'] or r['name'],r['spec'],r['unit'])
        groups.setdefault(key,[]).append(r)
        if r['code']: codes.setdefault((r.get('seller',''),r['code']),[]).append(r)
    for g in groups.values():
        rates={normalized_rate(v) for r in g for v in r['historical_rates']}
        for r in g:
            r['flags']=[]
            if not r['historical_rates']: r['flags'].append('缺税率')
            if len(rates)>1: r['flags'].append('历史有差异')
            if any(normalized_rate(v)=='0(含义待核)' for v in r['historical_rates']): r['flags'].append('0值待辨明')
            if not r['code']: r['flags'].append('缺商品编码')
            if any(c.get('type')=='f' for c in r.get('raw_cells',[])): r['flags'].append('原表含公式')
            if r['code'] and len(g)>1: r['flags'].append('商品编码重复')
            if not r['name']: r['flags'].append('缺商品名称')
    for g in codes.values():
        if len({(r['name'],r['spec'],r['unit']) for r in g})>1:
            for r in g: r['flags'].append('同码商品信息不同')
def import_rows(p):
    f,s,h,heads=sheet_info(p); m=p['mapping']; direction=p.get('direction','销项')
    if direction not in ['销项','进项']: raise ValueError('请选择进项或销项')
    if int(m.get('name',-1))<0: raise ValueError('必须选择商品名称列')
    selected=[int(v) for v in m.values() if int(v)>=0]
    if len(selected)!=len(set(selected)): raise ValueError('不同字段不能映射同一列，请核对。')
    if any(i>=len(heads) for i in selected): raise ValueError('映射列超出范围')
    rows=[]
    for rownum,raw in enumerate(s['rows'][h+1:],h+2):
        if all(c['value'] in [None,''] for c in raw): continue
        def get(k):
            i=int(m.get(k,-1)); return cell_text(raw[i]) if 0<=i<len(raw) else ''
        def val(k):
            i=int(m.get(k,-1)); return raw[i]['value'] if 0<=i<len(raw) else None
        dates=re.findall(r'\d{4}-\d{2}-\d{2}',get('date'))
        r={'id':uid(),'sample_id':'','code':get('code'),'name':get('name'),'spec':get('spec'),'unit':get('unit'),'seller':get('seller'),'dates':dates,'historical_rates':rate_values(val('rate')),'source_file':f['name'],'sheet':s['name'],'source_rows':[rownum],'raw_cells':raw,'history_examples':[],'status':'待复核','review':{'facts':get('facts'),'identity':get('identity')}}
        rows.append(r)
    if not rows: raise ValueError('没有可导入的数据行')
    b=fresh(f['name']+' · '+s['name'],rows)
    b['context']['direction']=direction
    sellers={r['seller'] for r in rows if r['seller']}
    if len(sellers)==1: b['context']['seller']=next(iter(sellers))
    b['source']={'token':p['token'],'headers':heads,'mapping':m,'sha256':f['sha256'],'header_row':h+1}
    enrich_flags(rows); persist(b); return view(b)
def sample():

    if not SAMPLE or not SAMPLE.is_file(): raise ValueError('未配置历史样本，请导入商品表开始复核。')
    source=read(SAMPLE)
    rows=[]
    for c in source['cases']:
        rows.append({'id':uid(),'sample_id':c['id'],'code':'','name':c['name'],'spec':c['spec'],'unit':c['unit'],'seller':c['seller'],'dates':c['transaction_dates'],'historical_rates':c['historical_rates'],'source_file':c['source_file'],'sheet':c['sheet'],'source_rows':c['source_rows'],'history_examples':c['history_examples'],'matching_record_count':c['matching_record_count'],'status':'待复核','review':{}})
    b=fresh('24组真实样本 · 未经财务确认',rows)
    b['context']['seller']='贵昌集团有限公司'; b['sample']=True
    b['source']={'sha256':source['source_sha256'],'note':'原开票底稿的24组候选，关联976条历史记录；TX编号不是ERP商品编码。'}
    enrich_flags(rows); persist(b); return view(b)
def dates_for(r,rev):
    return r['dates'] or [rev.get('transaction_date','')]
def validate_confirmation(b,r,rev):
    errors=[]; ctx=b['context']
    if not r['name']: errors.append('商品名称')
    for k,label in [('seller','销售主体'),('trade','主体、交易及优惠条件的核实记录'),('reviewer','复核人')]:
        if not ctx.get(k,'').strip(): errors.append(label)
    if ctx.get('taxpayer') in ['', '未核实']: errors.append('纳税人及计税身份')
    if r.get('seller') and r['seller']!=ctx.get('seller'): errors.append('原销售主体与本批主体不一致，请分批处理')
    for k,label in [('facts','商品事实'),('fact_source','事实来源'),('basis','确认理由与依据'),('valid_from','适用起日'),('valid_to','适用止日')]:
        if not str(rev.get(k,'')).strip(): errors.append(label)
    if rev.get('result') not in RATES: errors.append('明确的税务处理（数值0不作免税）')
    if rev.get('result')=='其他' and not rev.get('other_result','').strip(): errors.append('其他处理说明')
    if not rev.get('conditions_checked'): errors.append('勾选事实与交易条件已核对')
    dates=dates_for(r,rev)
    try:
        for x in dates+[rev.get('valid_from',''),rev.get('valid_to','')]: dt.date.fromisoformat(x)
        if not rev['valid_from']<=min(dates)<=max(dates)<=rev['valid_to']: errors.append('适用期间须覆盖本行全部交易日期')
        if (min(dates)<'2026-01-01' or max(dates)>'2026-09-20') and not rev.get('period_basis','').strip(): errors.append('内置核查期外的当期依据')
    except (ValueError,TypeError): errors.append('有效的交易日期和适用日期')
    if errors: raise ValueError('暂不能确认，请补齐：'+'；'.join(errors))
def save_context(p):
    b=get_batch(p['id']); check_rev(b,p)
    ctx={k:str(p['context'].get(k,'')).strip() for k in ['direction','seller','taxpayer','trade','reviewer']}
    if ctx['direction'] not in ['进项','销项']: raise ValueError('口径无效')
    if ctx['direction']!=b['context']['direction']: raise ValueError('本批进销项口径在导入时固定；请新导入并选择对应的进项或销项税率列。')
    if any(ctx[k]!=b['context'].get(k) for k in ['direction','seller','taxpayer','trade']):
        for r in b['rows']:
            if r['status']=='已确认': r['status']='待复核'; r['review']['conditions_checked']=False
    b['audit'].append({'at':now(),'action':'保存共同条件','before':b['context'],'after':ctx}); b['context']=ctx; persist(b); return view(b)
def save_review(p):
    b=get_batch(p['id']); check_rev(b,p)
    r=next(r for r in b['rows'] if r['id']==p['row_id'])
    if p['status'] not in ['待补充','已确认','待复核']: raise ValueError('状态无效')
    rev={k:str(p['review'].get(k,'')).strip() for k in ['result','other_result','facts','fact_source','basis','identity','valid_from','valid_to','transaction_date','period_basis']}
    rev['conditions_checked']=p['review'].get('conditions_checked') is True
    if p['status']=='已确认': validate_confirmation(b,r,rev)
    rev.update({'reviewer':b['context'].get('reviewer',''),'saved_at':now(),'policy_version':VERSION})
    b['audit'].append({'at':now(),'action':'商品复核','row_id':r['id'],'before':copy.deepcopy(r['review']),'before_status':r['status'],'after':rev,'status':p['status']})
    r['review']=rev; r['status']=p['status']; persist(b); return view(b)
def signature(b,r,rev):
    identity=r['code'] or rev.get('identity','').strip()
    if not identity or not rev.get('facts') or not rev.get('fact_source'): return None
    return (identity,r['name'],r['spec'],r['unit'],rev['facts'],rev['fact_source'],*(b['context'].get(k) for k in ['direction','seller','taxpayer','trade']),VERSION)
def reuse(p):
    b=get_batch(p['id']); r=next(r for r in b['rows'] if r['id']==p['row_id']); sig=signature(b,r,p['review']); matches=[]
    if not sig: return {'matches':[],'message':'需要商品编码或人工确认的商品身份，以及相同商品事实、事实来源、主体、进销项和交易条件。不会只凭名称复用。'}
    for path in (DATA/'batches').glob('*.json'):
        other=read(path)
        for x in other['rows']:
            if other['id']==b['id'] and x['id']==r['id']: continue
            rev=x['review']
            if x['status']!='已确认' or rev.get('policy_version')!=VERSION or sig!=signature(other,x,rev): continue
            dates=dates_for(r,p['review'])
            if not all(dates) or not rev['valid_from']<=min(dates)<=max(dates)<=rev['valid_to']: continue
            matches.append({'batch':other['title'],'name':x['name'],'review':rev})
    results={x['review']['result']+'|'+x['review'].get('other_result','') for x in matches}
    return {'matches':matches,'conflict':len(results)>1,'message':'找到条件一致的确认记录；填入后仍可修改并确认。' if matches else '没有找到条件、依据版本和适用期间均一致的已确认记录。'}
def compare(p):
    a,sa,ha,ca=sheet_info(p['before']); b,sb,hb,cb=sheet_info(p['after'])
    keys=p.get('keys',[])
    if not keys: raise ValueError('请选择用于对应两表的商品键；同名不视为唯一键。')
    if any(ca.count(k)!=1 or cb.count(k)!=1 for k in keys): raise ValueError('商品键列必须在两表中各出现一次')
    changes=[]; warnings=[]
    if ca!=cb: warnings.append('两表列名或列顺序不同')
    if [x['name'] for x in a['sheets']]!=[x['name'] for x in b['sheets']]: warnings.append('工作表名称或顺序不同')
    if len(sa['rows'])!=len(sb['rows']): warnings.append('两表总行数不同')
    if sa['merged']!=sb['merged']: warnings.append('合并单元格范围不同')
    if len(ca)!=len(set(ca)) or len(cb)!=len(set(cb)): warnings.append('有重复或空表头，不能可靠按列名对应全部列')
    def indexed(s,h,heads):
        idx={}; blanks=[]
        for n,row in enumerate(s['rows'][h+1:],h+2):
            if all(c['value'] in [None,''] for c in row): continue
            vals=[cell_text(row[heads.index(k)]) if heads.index(k)<len(row) else '' for k in keys]
            if not all(vals): blanks.append(n); continue
            key=tuple(vals); idx.setdefault(key,[]).append((n,row))
        return idx,blanks
    ia,ba=indexed(sa,ha,ca); ib,bb=indexed(sb,hb,cb)
    duplicates_a=[{'key':' / '.join(k),'rows':[x[0] for x in v]} for k,v in ia.items() if len(v)>1]
    duplicates_b=[{'key':' / '.join(k),'rows':[x[0] for x in v]} for k,v in ib.items() if len(v)>1]
    added=[' / '.join(k) for k in ib if k not in ia]; removed=[' / '.join(k) for k in ia if k not in ib]; moved=[]
    for k in ia.keys() & ib.keys():
        if len(ia[k])!=1 or len(ib[k])!=1: continue
        na,ra=ia[k][0]; nb,rb=ib[k][0]
        if na!=nb: moved.append({'key':' / '.join(k),'before_row':na,'after_row':nb})
        for col in ca:
            if not col or ca.count(col)!=1 or cb.count(col)!=1: continue
            c1=ra[ca.index(col)] if ca.index(col)<len(ra) else {'value':None,'type':'n','format':'General'}
            c2=rb[cb.index(col)] if cb.index(col)<len(rb) else {'value':None,'type':'n','format':'General'}
            if c1!=c2: changes.append({'key':' / '.join(k),'column':col,'before_row':na,'after_row':nb,'before':c1,'after':c2,'value_changed':c1['value']!=c2['value'],'type_changed':c1['type']!=c2['type'],'format_changed':c1['format']!=c2['format']})
    return {'before':{'file':a['name'],'sha256':a['sha256'],'rows':len(sa['rows']),'columns':ca},'after':{'file':b['name'],'sha256':b['sha256'],'rows':len(sb['rows']),'columns':cb},'warnings':warnings,'blank_keys_before':ba,'blank_keys_after':bb,'duplicates_before':duplicates_a,'duplicates_after':duplicates_b,'added_keys':added,'removed_keys':removed,'moved':moved,'changes':changes,'conclusion':'仅完成离线结构与值比较；缺少同次系统导入结果/报错，不能判定蔬东坡错乱根因，也不表示文件可直接回导。'}
def export_payload(b,mode):
    rows=b['rows'] if mode=='all' else [r for r in b['rows'] if r['status']=='已确认']
    if not rows: raise ValueError('还没有已确认商品，可先下载完整复核表。')
    headers=['样本编号','商品编码','商品名称','规格','单位','原税率原值','进销项','复核状态','确认处理','确认数值税率','销售主体','纳税/计税身份','交易条件','交易日期','适用起日','适用止日','商品事实','事实来源','确认理由及依据','候选建议','待补信息','官方条款定位','官方链接','来源文件','来源工作表','来源行号','复核人','确认时间','人工商品身份','依据版本','期外依据']
    values=[]; provenance=[]
    for r in rows:
        rev=r['review']; c=candidate(r,b['context']['direction']); result=rev.get('result','') if r['status']=='已确认' else ''
        numeric=float(result.rstrip('%'))/100 if result.endswith('%') else (0 if result=='零税率' else None)
        rowkey=r.get('sample_id') or r['id']
        for n in r['source_rows']: provenance.append([rowkey,r['code'],r['name'],r['source_file'],r['sheet'],n])
        transaction_dates=dates_for(r,rev)
        date_display=' / '.join(transaction_dates) if len(transaction_dates)<4 else f'{min(transaction_dates)} 至 {max(transaction_dates)}（{len(transaction_dates)}个日期；完整日期见JSON备份）'
        row_display=', '.join(map(str,r['source_rows'])) if len(r['source_rows'])<=8 else f'{len(r["source_rows"])}行；详见“来源行”表，组号{rowkey}'
        values.append([rowkey,r['code'],r['name'],r['spec'],r['unit'],' / '.join(r['historical_rates']),b['context']['direction'],r['status'],result+('：'+rev.get('other_result','') if result=='其他' else ''),numeric,b['context']['seller'],b['context']['taxpayer'],b['context']['trade'],date_display,rev.get('valid_from',''),rev.get('valid_to',''),rev.get('facts',''),rev.get('fact_source',''),rev.get('basis',''),c['summary'],c['missing'],c['locator'],'\n'.join(SOURCES[k]['url'] for k in c['sources']),r['source_file'],r['sheet'],row_display,rev.get('reviewer','') if result else '',export_time(rev.get('saved_at','')) if result else '',rev.get('identity',''),VERSION,rev.get('period_basis','')])
    return {'title':b['title'],'note':'复核结果表，非ERP导入模板。免税的数值税率留空，零税率才写数值0。未确认行不填确认处理。','headers':headers,'rows':values,'sheet':'税率复核','percent_columns':[9],'provenance':provenance}
def export_xlsx(p):
    b=get_batch(p['id']); check_rev(b,p); payload=export_payload(b,p.get('mode','confirmed')); token=uid()
    inp=identified('exports',token); out=identified('exports',token,'.xlsx'); dump(inp,payload)
    write_workbook(payload,out)
    return {'url':'/download/'+token+'.xlsx','rows':len(payload['rows']),'name':('全部复核' if p.get('mode')=='all' else '已确认结果')+'.xlsx'}

class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args): pass
    def send(self,value,status=200):
        raw=json.dumps(value,ensure_ascii=False).encode(); self.send_response(status); self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(raw)
    def local_request(self):
        if self.headers.get('Host') not in [f'127.0.0.1:{PORT}',f'localhost:{PORT}']: raise ValueError('仅允许本机访问')
        origin=self.headers.get('Origin')
        if origin and origin not in [f'http://127.0.0.1:{PORT}',f'http://localhost:{PORT}']: raise ValueError('请从本机Demo页面操作')
    def do_GET(self):
        try:
            self.local_request(); path=urlparse(self.path); q=parse_qs(path.query)
            if path.path.startswith('/api/v2/'):
                return self.send(workflow().get(path.path[len('/api/v2/'):],{k:v[0] for k,v in q.items()}))
            if path.path.startswith('/download/v2/'):
                name=path.path.split('/')[-1]
                if not re.fullmatch(r'[a-f0-9]{32}\.(xlsx|csv|json)',name):raise ValueError('导出编号无效')
                file=DATA/'v2'/name
                self.send_response(200);self.send_header('Content-Type','application/octet-stream');self.send_header('Content-Disposition','attachment; filename='+name);self.send_header('Cache-Control','no-store');self.end_headers();self.wfile.write(file.read_bytes());return
            if path.path=='/api/batches':
                files=[read(p) for p in (DATA/'batches').glob('*.json')]
                return self.send({'batches':[{'id':b['id'],'title':b['title'],'saved_at':b['saved_at'],'count':len(b['rows']),'confirmed':sum(r['status']=='已确认' for r in b['rows'])} for b in sorted(files,key=lambda b:b['saved_at'],reverse=True)],'sources':SOURCES,'version':VERSION})
            if path.path=='/api/batch': return self.send(view(get_batch(q['id'][0])))
            if path.path=='/api/backup':
                b=get_batch(q['id'][0]); self.send_response(200); self.send_header('Content-Type','application/json'); self.send_header('Content-Disposition',"attachment; filename=review-backup.json"); self.end_headers(); self.wfile.write(json.dumps(b,ensure_ascii=False,indent=2).encode()); return
            if path.path.startswith('/download/'):
                token=path.path.split('/')[-1].removesuffix('.xlsx'); file=identified('exports',token,'.xlsx'); content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            else:
                name=path.path.lstrip('/') or 'index.html'
                if name not in ['index.html','legacy.html','app.js','style.css','v2.js','v2.css']: return self.send({'error':'未找到'},404)
                file=ROOT/'static'/name; content_type={'.html':'text/html; charset=utf-8','.js':'text/javascript; charset=utf-8','.css':'text/css; charset=utf-8'}[file.suffix]
            raw=file.read_bytes(); self.send_response(200); self.send_header('Content-Type',content_type); self.send_header('Cache-Control','no-store'); self.send_header('X-Content-Type-Options','nosniff'); self.end_headers(); self.wfile.write(raw)
        except FileNotFoundError: self.send({'error':'未找到当前账号的文件或批次'},404)
        except (ValueError,KeyError) as e: self.send({'error':str(e)},400)
    def do_POST(self):
        try:
            self.local_request(); length=int(self.headers.get('Content-Length','0'))
            if not 0<length<=45*1024*1024: raise ValueError('文件或请求过大')
            p=json.loads(self.rfile.read(length)); path=urlparse(self.path).path
            if path.startswith('/api/v2/'):return self.send(workflow().post(path[len('/api/v2/'):],p))
            with LOCK:
                if path=='/api/sample': result=sample()
                elif path=='/api/upload': result=upload(p)
                elif path=='/api/preview': result=preview(p)
                elif path=='/api/import': result=import_rows(p)
                elif path=='/api/context': result=save_context(p)
                elif path=='/api/review': result=save_review(p)
                elif path=='/api/reuse': result=reuse(p)
                elif path=='/api/compare': result=compare(p)
                elif path=='/api/export': result=export_xlsx(p)
                elif path=='/api/restore':
                    b=p['batch']
                    if b.get('schema')!=1 or not isinstance(b.get('rows'),list): raise ValueError('不是本Demo的备份')
                    b['id']=uid(); b['revision']=0; b['title']+=' · 恢复副本'
                    for r in b['rows']:
                        r['id']=uid()
                        if r['status']=='已确认':
                            validate_confirmation(b,r,r['review'])
                            if r['review'].get('policy_version')!=VERSION: r['status']='待复核'
                    persist(b); result=view(b)
                else: raise ValueError('操作不存在')
            self.send(result)
        except Exception as e: self.send({'error':str(e)},400)

if __name__=='__main__':
    for folder in ['batches','uploads','exports']: (DATA/folder).mkdir(parents=True,exist_ok=True)
    print(f'税率复核 Demo http://127.0.0.1:{PORT}  数据：{DATA}',flush=True)
    ThreadingHTTPServer(('127.0.0.1',PORT),Handler).serve_forever()
