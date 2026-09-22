"""Exact purchase-number search with source locations; imports are never payment proof."""
import csv, io, re, unicodedata, zipfile
from pathlib import Path

def extract(name, raw):
    ext=Path(name).suffix.lower(); records=[]; gaps=[]
    def add(location,value):
        if value is not None and str(value).strip():records.append({'location':location,'text':str(value)[:10000]})
    if ext in ['.xlsx','.xlsm']:
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>80_000_000:raise ValueError('Excel解压后过大，请缩小范围')
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            for ws in wb:
                if ws.max_row>50000 or ws.max_column>300:raise ValueError('表格范围过大，请裁剪空白格式区域')
                for row in ws:
                    for c in row:
                        if c.value is not None:add(f'{ws.title}!{c.coordinate}',c.value)
        finally:wb.close()
    elif ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw))
        if len(reader.pages)>300:raise ValueError('PDF最多300页')
        for i,page in enumerate(reader.pages,1):
            text=page.extract_text() or ''
            if not text.strip():gaps.append(f'第{i}页无可提取文字，需OCR或人工核对')
            for j,line in enumerate(text.splitlines(),1):add(f'第{i}页 第{j}行',line)
    elif ext in ['.csv','.tsv','.txt']:
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:text=raw.decode('gb18030')
        if ext=='.txt':
            for i,l in enumerate(text.splitlines(),1):add(f'第{i}行',l)
        else:
            for i,row in enumerate(csv.reader(io.StringIO(text),delimiter='\t' if ext=='.tsv' else ','),1):
                for j,v in enumerate(row,1):add(f'第{i}行 第{j}列',v)
    else:raise ValueError('支持XLSX、CSV、TSV、TXT和文字PDF；旧XLS或图片请先另存/转文字')
    if len(records)>150000:raise ValueError('可检索内容过多，请分批导入')
    if not records and not gaps:gaps.append('附件未提取到文本')
    return records,gaps

def normalize(s):return unicodedata.normalize('NFKC',s).upper().strip()
def search(docs,query):
    q=normalize(query)
    if not re.fullmatch(r'[A-Z0-9][A-Z0-9_-]{2,99}',q):raise ValueError('请输入完整采购单号（至少3位字母、数字、-或_）')
    pattern=re.compile(r'(?<![A-Z0-9_-])'+re.escape(q)+r'(?![A-Z0-9_-])')
    hits=[]
    for d in docs:
        for r in d['records']:
            if pattern.search(normalize(r['text'])):
                hits.append({**{k:d[k] for k in ['id','filename','payment_no','subject','period','sha256']},**r})
    return dict(query=q,hits=hits,files_searched=len(docs),unreadable_sections=sum(len(d['gaps']) for d in docs),
      periods=sorted(set(d['period'] for d in docs)),
      conclusion='发现附件文字命中，需核付款状态及是否漏点结算' if hits else '在已导入范围内未命中，不能直接认定未报销',
      pending=['核对付款单是否审批通过、实际支付','核对采购单是否退货/撤销','确认导入主体、日期和人员范围完整'] )
